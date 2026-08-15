import os
import re
import sys
import socket
import atexit
import logging
import threading
from concurrent.futures import ThreadPoolExecutor, as_completed

import pandas as pd
from openpyxl import load_workbook
from openpyxl.drawing.image import Image
from decimal import Decimal, ROUND_DOWN

import json
import queue
import io
import base64
import hashlib
from datetime import datetime
import subprocess
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, simpledialog, font as tkfont
from PIL import Image as PilImage, ImageTk, ImageDraw


# ======================== OCR 管道引擎 ========================

class PPOCR_pipe:
    """调用 PaddleOCR-json.exe 的 Python Api（管道模式）"""

    def __init__(self, exePath: str, modelsPath: str = None, argument: dict = None):
        self.__ENABLE_CLIPBOARD = False

        exePath = os.path.abspath(exePath)
        cwd = os.path.abspath(os.path.join(exePath, os.pardir))
        cmds = [exePath]
        if modelsPath is not None:
            if os.path.exists(modelsPath) and os.path.isdir(modelsPath):
                cmds += ["--models_path", os.path.abspath(modelsPath)]
            else:
                raise Exception(
                    f"Input modelsPath doesn't exist or isn't a directory. modelsPath: [{modelsPath}]"
                )
        if isinstance(argument, dict):
            for key, value in argument.items():
                if isinstance(value, bool):
                    cmds += [f"--{key}={value}"]
                elif isinstance(value, str):
                    cmds += [f"--{key}", value]
                else:
                    cmds += [f"--{key}", str(value)]
        self.ret = None
        startupinfo = None
        if "win32" in str(sys.platform).lower():
            startupinfo = subprocess.STARTUPINFO()
            startupinfo.dwFlags = (
                subprocess.CREATE_NO_WINDOW | subprocess.STARTF_USESHOWWINDOW
            )
            startupinfo.wShowWindow = subprocess.SW_HIDE
        self.ret = subprocess.Popen(
            cmds,
            cwd=cwd,
            stdin=subprocess.PIPE,
            stdout=subprocess.PIPE,
            stderr=subprocess.DEVNULL,
            startupinfo=startupinfo,
        )
        while True:
            if not self.ret.poll() is None:
                raise Exception(f"OCR init fail.")
            initStr = self.ret.stdout.readline().decode("utf-8", errors="ignore")
            if "OCR init completed." in initStr:
                break
            elif "OCR clipboard enabled." in initStr:
                self.__ENABLE_CLIPBOARD = True
        atexit.register(self.exit)

    def runDict(self, writeDict: dict):
        if not self.ret:
            return {"code": 901, "data": f"引擎实例不存在。"}
        if not self.ret.poll() is None:
            return {"code": 902, "data": f"子进程已崩溃。"}
        writeStr = json.dumps(writeDict, ensure_ascii=True, indent=None) + "\n"
        try:
            self.ret.stdin.write(writeStr.encode("utf-8"))
            self.ret.stdin.flush()
        except Exception as e:
            return {
                "code": 902,
                "data": f"向识别器进程传入指令失败，疑似子进程已崩溃。{e}",
            }
        try:
            getStr = self.ret.stdout.readline().decode("utf-8", errors="ignore")
        except Exception as e:
            return {"code": 903, "data": f"读取识别器进程输出值失败。异常信息：[{e}]"}
        try:
            return json.loads(getStr)
        except Exception as e:
            return {
                "code": 904,
                "data": f"识别器输出值反序列化JSON失败。异常信息：[{e}]。原始内容：[{getStr}]",
            }

    def run(self, imgPath: str):
        return self.runDict({"image_path": imgPath})

    def runClipboard(self):
        if self.__ENABLE_CLIPBOARD:
            return self.run("clipboard")
        else:
            raise Exception("剪贴板功能不存在或已禁用。")

    def runBase64(self, imageBase64: str):
        return self.runDict({"image_base64": imageBase64})

    def runBytes(self, imageBytes):
        imageBase64 = base64.b64encode(imageBytes).decode("utf-8")
        return self.runBase64(imageBase64)

    def exit(self):
        if hasattr(self, "ret"):
            if not self.ret:
                return
            try:
                self.ret.kill()
            except Exception as e:
                print(f"[Error] ret.kill() {e}")
        self.ret = None
        atexit.unregister(self.exit)

    @staticmethod
    def printResult(res: dict):
        if res["code"] == 100:
            index = 1
            for line in res["data"]:
                print(
                    f"{index}-置信度：{round(line['score'], 2)}，文本：{line['text']}"
                )
                index += 1
        elif res["code"] == 101:
            print("图片中未识别出文字。")
        else:
            print(f"图片识别失败。错误码：{res['code']}，错误信息：{res['data']}")

    def __del__(self):
        self.exit()

# ======================== RSA 工具函数 ========================

def _der_read_length(data, offset):
    """读取 ASN.1 DER 长度"""
    b = data[offset]
    if b < 0x80:
        return b, offset + 1
    num_bytes = b & 0x7F
    length = int.from_bytes(data[offset + 1:offset + 1 + num_bytes], 'big')
    return length, offset + 1 + num_bytes


def _der_read_integer(data, offset):
    """读取 ASN.1 DER INTEGER"""
    assert data[offset] == 0x02, f"Expected INTEGER tag, got {data[offset]:#x}"
    length, offset = _der_read_length(data, offset + 1)
    value = int.from_bytes(data[offset:offset + length], 'big', signed=False)
    return value, offset + length


def _der_read_sequence(data, offset):
    """读取 ASN.1 SEQUENCE, 返回 (content_bytes, new_offset)"""
    assert data[offset] == 0x30, f"Expected SEQUENCE tag, got {data[offset]:#x}"
    length, offset = _der_read_length(data, offset + 1)
    return data[offset:offset + length], offset + length


def _pem_to_der(pem_data):
    """PEM 转 DER"""
    lines = pem_data.strip().splitlines()
    b64_lines = [l for l in lines if not l.startswith('-----')]
    return base64.b64decode(''.join(b64_lines))


def _parse_private_key_pem(pem_path):
    """解析 PKCS#8 PrivateKeyInfo PEM, 返回 (n, e, d)"""
    with open(pem_path, 'rb') as f:
        pem_data = f.read().decode('ascii')
    der = _pem_to_der(pem_data)
    pki, _ = _der_read_sequence(der, 0)
    offset = 0
    # version
    _, offset = _der_read_integer(pki, offset)
    # algorithm SEQUENCE
    _, offset = _der_read_sequence(pki, offset)
    # privateKey OCTET STRING
    assert pki[offset] == 0x04, f"Expected OCTET STRING, got {pki[offset]:#x}"
    length, offset = _der_read_length(pki, offset + 1)
    rsa_der = pki[offset:offset + length]
    # RSAPrivateKey SEQUENCE { version, n, e, d, ... }
    rsa_priv, _ = _der_read_sequence(rsa_der, 0)
    off = 0
    _, off = _der_read_integer(rsa_priv, off)
    n, off = _der_read_integer(rsa_priv, off)
    e, off = _der_read_integer(rsa_priv, off)
    d, _ = _der_read_integer(rsa_priv, off)
    return n, e, d


# SHA-256 OID 编码后的 DigestInfo 前缀
_SHA256_DIGEST_INFO_PREFIX = bytes([
    0x30, 0x31, 0x30, 0x0d, 0x06, 0x09, 0x60, 0x86, 0x48, 0x01,
    0x65, 0x03, 0x04, 0x02, 0x01, 0x05, 0x00, 0x04, 0x20
])


def rsa_sign(message, private_key_path):
    """PKCS#1 v1.5 RSA-SHA256 签名, 返回 base64 字符串"""
    n, e, d = _parse_private_key_pem(private_key_path)
    k = (n.bit_length() + 7) // 8

    digest = hashlib.sha256(message.encode('utf-8')).digest()
    digest_info = _SHA256_DIGEST_INFO_PREFIX + digest

    ps_len = k - len(digest_info) - 3
    if ps_len < 8:
        raise ValueError("密钥长度不足")
    padded = b'\x00\x01' + b'\xff' * ps_len + b'\x00' + digest_info

    m = int.from_bytes(padded, 'big')
    s = pow(m, d, n)
    return base64.b64encode(s.to_bytes(k, 'big')).decode('ascii')


def get_machine_id():
    """获取本机机器码（仅主板序列号，不包含 MAC 地址）
    优先读取 launcher 预设的环境变量，避免 Cython 编译后 subprocess 异常。"""
    mid = os.environ.get('_MACHINE_ID')
    if mid:
        return mid
    sn = ''
    # 方法1: wmic（旧版 Windows）
    try:
        result = subprocess.run(
            ['wmic', 'baseboard', 'get', 'serialnumber'],
            capture_output=True, text=True, timeout=5,
            creationflags=subprocess.CREATE_NO_WINDOW
        )
        lines = [l.strip() for l in result.stdout.strip().splitlines() if l.strip()]
        if len(lines) > 1 and lines[-1] and not lines[-1].lower().startswith('serial'):
            sn = lines[-1]
    except Exception:
        pass
    # 方法2: PowerShell（Windows 10/11 推荐）
    if not sn:
        try:
            result = subprocess.run(
                ['powershell', '-Command',
                 '(Get-CimInstance Win32_BaseBoard).SerialNumber'],
                capture_output=True, text=True, timeout=5,
                creationflags=subprocess.CREATE_NO_WINDOW
            )
            sn = result.stdout.strip()
        except Exception:
            pass
    return sn


# ======================== License 验证 ========================

_PUBKEY_N = 23226204762270336453441542784267745136506622375676152815303942589897687286985912735863376575686137348348998792058972048729168654274848222713272035275183713068836087331347801504893443785325656862408749898399063311718679629024725481870672842363844245424236052362957785654737771332697971035683223897974618569110215559562573439182578191744664534975787322724684855839154114378379947466580881843608915507170898557800406027107379707859546718279612369202517721923479859814790372173551656305072687854980323214582073787946312828885136371243325846362275777949021968770880746288390861032238105467806987526222491888828925394751113
_PUBKEY_E = 65537


def _rsa_verify_embedded(message, signature_b64):
    """用内置公钥验证签名，不需要 public_key.pem 文件"""
    k = (_PUBKEY_N.bit_length() + 7) // 8

    try:
        sig_bytes = base64.b64decode(signature_b64)
    except Exception:
        return False
    if len(sig_bytes) != k:
        return False

    s = int.from_bytes(sig_bytes, 'big')
    m = pow(s, _PUBKEY_E, _PUBKEY_N)
    decrypted = m.to_bytes(k, 'big')

    if decrypted[0:2] != b'\x00\x01':
        return False
    try:
        sep_idx = decrypted.index(b'\x00', 2)
    except ValueError:
        return False
    if sep_idx - 2 < 8:
        return False

    digest_info = decrypted[sep_idx + 1:]
    expected_digest = hashlib.sha256(message.encode('utf-8')).digest()
    expected_di = _SHA256_DIGEST_INFO_PREFIX + expected_digest
    return digest_info == expected_di


def _show_license_error(title, message):
    """显示授权失败弹窗，附带机器码展示"""
    mid = get_machine_id()

    dialog = tk.Tk()
    dialog.title(title)
    dialog.resizable(False, False)

    # 居中
    w, h = 440, 280
    sw = dialog.winfo_screenwidth()
    sh = dialog.winfo_screenheight()
    dialog.geometry(f"{w}x{h}+{(sw - w) // 2}+{(sh - h) // 2}")

    frm = ttk.Frame(dialog, padding=16)
    frm.pack(fill="both", expand=True)

    ttk.Label(frm, text="⚠", font=("", 36)).pack(pady=(0, 4))
    ttk.Label(frm, text=message, font=("", 10), wraplength=400,
              justify="center").pack()

    mid_frame = ttk.LabelFrame(frm, text="本机机器码", padding=6)
    mid_frame.pack(fill="x", pady=12)

    mid_var = tk.StringVar(value=mid)
    ttk.Entry(mid_frame, textvariable=mid_var, font=("Consolas", 13),
              justify="center", state="readonly").pack(fill="x", padx=4, pady=2)

    # 按钮 + 作者信息
    bottom = ttk.Frame(frm)
    bottom.pack(fill="x", pady=(12, 0))

    ttk.Label(bottom, text="", foreground="gray").pack(side="left")
    ttk.Button(bottom, text="退出", command=dialog.destroy).pack(side="right")

    dialog.protocol("WM_DELETE_WINDOW", dialog.destroy)
    dialog.mainloop()
    sys.exit(1)


def _get_app_dir():
    """获取应用程序目录（exe 所在目录 或 脚本所在目录）"""
    if getattr(sys, 'frozen', False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


def check_license():
    """验证 license.lic，失败则弹窗退出。返回 True 表示通过"""
    script_dir = _get_app_dir()
    license_path = os.path.join(script_dir, "license.lic")

    if not os.path.exists(license_path):
        _show_license_error(
            "授权验证失败",
            "未找到 license.lic 授权文件。\n请将本机机器码发给软件提供方获取授权。"
        )

    try:
        with open(license_path, 'r', encoding='utf-8') as f:
            lic = json.load(f)
        data = lic.get("data", {})
        signature = lic.get("signature", "")
        message = json.dumps(data, ensure_ascii=False, separators=(',', ':'))
        if not _rsa_verify_embedded(message, signature):
            raise ValueError("签名验证失败")
    except Exception as e:
        _show_license_error(
            "授权验证失败",
            f"License 文件无效或已被篡改。\n请将本机机器码发给软件提供方获取新授权。"
        )

    if not isinstance(data, dict):
        _show_license_error(
            "授权验证失败",
            "License 文件内容无效。\n请将本机机器码发给软件提供方获取授权。"
        )

    # 检查机器码
    machine_id = data.get("machine_id", "")
    if machine_id and machine_id != "any":
        if machine_id != get_machine_id():
            _show_license_error(
                "授权验证失败",
                f"本机机器码与授权文件不匹配。\n请将下方的机器码发给软件提供方。"
            )

    # 检查有效期 + 防修改系统时间
    expire_date = data.get("expire_date", "")
    if expire_date and expire_date != "never":
        try:
            exp = datetime.strptime(expire_date, "%Y-%m-%d")
            now = datetime.now()
            if now > exp:
                _show_license_error(
                    "授权已过期",
                    f"License 已于 {expire_date} 过期。\n请联系软件提供方续期。"
                )
            _check_clock_rollback(script_dir, now)
        except ValueError:
            pass

    return True


# ---- 防修改系统时间 ----

_RT_NAME = ".runtime"


def _get_runtime_path(script_dir):
    return os.path.join(script_dir, _RT_NAME)


def _check_clock_rollback(script_dir, now):
    """检测系统时间是否被回调，是则拒绝运行"""
    path = _get_runtime_path(script_dir)
    ts = int(now.timestamp())

    if os.path.exists(path):
        try:
            with open(path, 'rb') as f:
                saved_bytes = f.read()
            saved = int.from_bytes(saved_bytes, 'big') ^ 0xA3E7F1
            if saved > ts:
                _show_license_error(
                    "授权验证失败",
                    "检测到系统时间异常，请恢复正确时间后重试。"
                )
        except Exception:
            pass

    try:
        with open(path, 'wb') as f:
            f.write((ts ^ 0xA3E7F1).to_bytes(8, 'big'))
    except Exception as e:
        print(f"写入时间记录失败: {e}")


def _find_paddleocr_dir():
    """查找 PaddleOCR-json 引擎目录（支持打包、脚本运行和环境变量）"""
    env_dir = os.environ.get("PADDLEOCR_DIR", "").strip()
    if env_dir and os.path.isfile(os.path.join(env_dir, "PaddleOCR-json.exe")):
        return env_dir

    candidates = []
    if getattr(sys, 'frozen', False):
        candidates.append(os.path.join(sys._MEIPASS, "PaddleOCR-json"))
        candidates.append(os.path.join(os.path.dirname(sys.executable), "PaddleOCR-json"))
    candidates.append(os.path.join(os.path.dirname(os.path.abspath(__file__)), "PaddleOCR-json"))
    candidates.append(os.path.join(os.getcwd(), "PaddleOCR-json"))

    for cand in candidates:
        if os.path.isfile(os.path.join(cand, "PaddleOCR-json.exe")):
            return cand
    return candidates[0] if candidates else "PaddleOCR-json"


# ======================== 统一路径配置 ========================
# 修改这里即可适配不同环境
BASE_DIR = r"D:\baogao\sferzaibo"
PING_DIR = os.path.join(BASE_DIR, "ping")
PICTUER_DIR = os.path.join(BASE_DIR, "pictuer")
JIETU_DIR = os.path.join(BASE_DIR, "jietu")
DT_DIR = os.path.join(BASE_DIR, "dt")
RESULT_EXCEL = os.path.join(BASE_DIR, "result.xlsx")
FINISH_EXCEL = os.path.join(BASE_DIR, "finish.xlsx")
TEMPLATE_EXCEL = os.path.join(BASE_DIR, "室分模板.xlsx")

# OCR 引擎路径（自动查找，兼容 exe 打包、脚本运行和环境变量）
PADDLEOCR_DIR = _find_paddleocr_dir()

SCRIPT_DIR = _get_app_dir()
PROFILES_FILE = os.path.join(SCRIPT_DIR, "gui_profiles.json")



# ======================== 工具函数 ========================

def ensure_dir(path):
    """确保目录存在"""
    if not os.path.exists(path):
        os.makedirs(path)


def is_image(filename):
    return filename.lower().endswith(('.jpg', '.png', '.jpeg', '.bmp'))


def _open_template(path):
    """打开模板文件，失败时给出清晰提示"""
    try:
        return load_workbook(path)
    except Exception as e:
        if not os.path.exists(path):
            raise FileNotFoundError(f"模板文件不存在: {path}")
        raise type(e)(
            f"无法打开模板文件 {path}，请确认它是 .xlsx 格式（非 .xls），且未损坏。\n原始错误: {e}"
        )


def _insert_image(ws, img, cell, idx):
    """插入图片到指定单元格，使用 OneCellAnchor。

    图片属性 = "随单元格改变位置，但不改变大小"（Excel 中
    "大小和属性 -> 属性 -> 移动但不随单元格调整大小"）：
    - 位置 = 单元格坐标（随单元格移动）
    - 大小 = 绝对 EMU 尺寸（ext，不随单元格变化）
    与字体/列宽换算完全无关，任何端打开都显示完全相同的大小
    （= 配置尺寸，缩放比例 100%）。
    """
    from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor
    from openpyxl.drawing.picture import PictureFrame
    from openpyxl.drawing.fill import Blip
    from openpyxl.drawing.geometry import PresetGeometry2D
    from openpyxl.utils.units import pixels_to_EMU
    from openpyxl.utils import coordinate_to_tuple

    row, col = coordinate_to_tuple(str(cell).upper())
    anchor = OneCellAnchor()
    anchor._from.row = row - 1
    anchor._from.col = col - 1
    anchor.ext.width = pixels_to_EMU(img.width)
    anchor.ext.height = pixels_to_EMU(img.height)

    pic = PictureFrame()
    pic.nvPicPr.cNvPr.id = idx
    pic.nvPicPr.cNvPr.name = "Image {0}".format(idx)
    # 与 Excel 原生一致，不写 descr/cstate（openpyxl 默认值不影响渲染）
    pic.nvPicPr.cNvPr.descr = None
    pic.blipFill.blip = Blip()
    pic.blipFill.blip.embed = "rId{0}".format(idx)   # 保存时会自动修正
    pic.spPr.prstGeom = PresetGeometry2D(prst="rect")
    anchor.pic = pic

    img.anchor = anchor
    ws.add_image(img)


def _resize_image_to_cm(img_path, width_cm, height_cm, dpi=96):
    """物理缩放图片到指定厘米尺寸（默认 96 DPI）。

    关键：让"图片文件本身的像素尺寸"就等于"期望的显示尺寸"。
    这样无论渲染器按 ext 属性还是按图片原始像素渲染，
    电脑/手机/微信/别人电脑看到的图片大小都完全一致，
    避免出现"电脑正常、手机空白、别人电脑超出"的差异。
    返回 (BytesIO, 宽px, 高px)。
    """
    px_per_cm = dpi / 2.54
    target_w = max(1, int(round(width_cm * px_per_cm)))
    target_h = max(1, int(round(height_cm * px_per_cm)))
    with PilImage.open(img_path) as pil:
        resized = pil.resize((target_w, target_h), PilImage.LANCZOS)
        buf = io.BytesIO()
        resized.save(buf, format="PNG", dpi=(dpi, dpi))
    buf.seek(0)
    return buf, target_w, target_h


def insert_images_to_excel(wb, sheet_name, image_positions, clear_first=True):
    """
    将图片批量插入到 Excel 工作表的指定位置。
    严格按配置的 width_cm / height_cm 插入：物理缩放图片像素到该尺寸，
    并用 OneCellAnchor 绝对尺寸锚定（位置 = 单元格坐标，大小 = 绝对 EMU），
    与字体/列宽换算无关，任何端打开都显示完全相同的大小。
    image_positions: {img_path: {"position": "A1", "width_cm": 10, "height_cm": 8}}
    """
    ws = wb[sheet_name]
    if clear_first:
        ws._images.clear()

    for idx, (img_path, cfg) in enumerate(image_positions.items(), 1):
        if not os.path.exists(img_path):
            print(f"警告: 图片 {img_path} 不存在，跳过。")
            continue
        # 物理缩放图片像素 = 配置尺寸（原始尺寸 = 配置，缩放比例 100%）
        buf, w_px, h_px = _resize_image_to_cm(
            img_path, cfg["width_cm"], cfg["height_cm"])
        img = Image(buf)
        img.width = w_px
        img.height = h_px
        _insert_image(ws, img, cfg["position"], idx)


def clear_folder(folder_path):
    """清空文件夹中的所有文件"""
    if not os.path.isdir(folder_path):
        print(f"目录不存在，跳过清理: {folder_path}")
        return
    for filename in os.listdir(folder_path):
        file_path = os.path.join(folder_path, filename)
        try:
            if os.path.isfile(file_path):
                os.remove(file_path)
                print(f"已删除文件: {file_path}")
        except Exception as e:
            print(f"删除文件 {file_path} 失败: {e}")


# ======================== OCR 引擎 ========================

# 并发识别时最多启动的独立引擎实例数（每个实例约占用 300~600MB 内存，
# 请根据机器内存调整；设为 1 则退化为串行）
_OCR_CONCURRENCY = 3

_ocr_engine = None
_ocr_lock = threading.Lock()
_ocr_local = threading.local()   # 每个工作线程一个独立引擎，实现真正并发


def _create_ocr_engine():
    """创建并启动一个 PaddleOCR-json 引擎实例"""
    logging.info("正在启动 PaddleOCR-json 引擎...")
    engine = PPOCR_pipe(
        exePath=os.path.join(PADDLEOCR_DIR, "PaddleOCR-json.exe"),
        modelsPath=os.path.join(PADDLEOCR_DIR, "models"),
        argument={
            "config_path": "models/config_chinese.txt",
            "cls": False,
            "limit_side_len": 1024,
            "enable_mkldnn": True,
        }
    )
    logging.info("PaddleOCR-json 引擎启动完成")
    return engine


def _get_ocr_engine():
    """获取全局单例引擎（供串行/非并发场景使用）"""
    global _ocr_engine
    if _ocr_engine is None:
        with _ocr_lock:
            if _ocr_engine is None:
                _ocr_engine = _create_ocr_engine()
    return _ocr_engine


def _get_thread_ocr_engine():
    """获取当前线程独立的引擎实例（多线程并发使用，互不阻塞）"""
    engine = getattr(_ocr_local, "engine", None)
    if engine is None:
        engine = _create_ocr_engine()
        _ocr_local.engine = engine
    return engine


def ocr_image(image_path, thread_engine=False):
    """使用 PaddleOCR-json 引擎识别图片文字（管道模式）
    thread_engine=True 时使用当前线程独立的引擎实例，可多线程并发识别"""
    try:
        if thread_engine:
            # 线程独立引擎，无需加锁
            res = _get_thread_ocr_engine().run(image_path)
        else:
            engine = _get_ocr_engine()
            with _ocr_lock:
                res = engine.run(image_path)
        code = res.get("code")
        if code == 100:
            data = res.get("data", [])
            if data:
                return " ".join(item.get("text", "") for item in data)
            return "未识别到文本"
        elif code == 101:
            return "未识别到文本"
        else:
            return f"OCR 失败: {res}"
    except Exception as e:
        logging.error(f"OCR 识别失败: {image_path}, 错误: {e}")
        return f"OCR 出错: {e}"


# ======================== 第一步：裁剪 ping 图片 ========================

def _import_cv2():
    """导入 cv2，自动注册 DLL 搜索路径"""
    import importlib, importlib.util, os
    spec = importlib.util.find_spec("cv2")
    if spec and spec.origin:
        cv2_dir = os.path.dirname(spec.origin)
        os.add_dll_directory(cv2_dir)
    return importlib.import_module("cv2")


def _cv2_imread(cv2, path):
    """支持中文路径的图片读取"""
    import numpy as np
    data = np.fromfile(path, dtype=np.uint8)
    if data.size == 0:
        return None
    return cv2.imdecode(data, cv2.IMREAD_COLOR)


def _cv2_imwrite(cv2, path, image):
    """支持中文路径的图片保存"""
    import numpy as np
    ext = os.path.splitext(path)[1] or ".jpg"
    ok, buf = cv2.imencode(ext, image)
    if ok:
        buf.tofile(path)
    return ok


def step1_crop_ping(input_folder, output_folder, regions=None):
    """从 ping 截图中裁剪指定区域
    regions: [(name, x1, y1, x2, y2), ...]  默认 delay + jitter"""
    cv2 = _import_cv2()
    if regions is None:
        regions = [("delay", 572, 30, 639, 51), ("jitter", 506, 79, 567, 95)]
    ensure_dir(output_folder)

    for filename in os.listdir(input_folder):
        if not is_image(filename):
            continue
        image = _cv2_imread(cv2, os.path.join(input_folder, filename))
        if image is None:
            print(f"无法加载图片: {filename}")
            continue

        base_name = os.path.splitext(filename)[0]
        h, w = image.shape[:2]
        for name, x1, y1, x2, y2 in regions:
            x1c, y1c = max(x1, 0), max(y1, 0)
            x2c, y2c = min(x2, w), min(y2, h)
            if x2c <= x1c or y2c <= y1c:
                print(f"警告: {filename} 的区域 {name} 坐标无效或越界，跳过。")
                continue
            crop = image[y1c:y2c, x1c:x2c]
            out_path = os.path.join(output_folder, f'{base_name}_{name}.jpg')
            if not _cv2_imwrite(cv2, out_path, crop):
                print(f"警告: 保存图片失败 {out_path}")
        print(f"处理并保存图片: {filename}")

    print("第一步：ping 截图裁剪完成")


# ======================== 第二步：裁剪主数据图片 ========================

def step2_crop_main(input_folder, output_folder, regions=None):
    """从主截图中裁剪指定点位区域
    regions: [(name, x1, y1, x2, y2), ...]  默认 8 个点位"""
    cv2 = _import_cv2()
    if regions is None:
        regions = [
            ("pindian_1", 137, 229, 198, 251), ("rsrp_1", 449, 29, 512, 53),
            ("sinr_1", 451, 55, 513, 77), ("pindian_2", 226, 229, 283, 251),
            ("rsrp_2", 516, 29, 578, 52), ("sinr_2", 516, 55, 579, 77),
            ("dl", 141, 648, 276, 675), ("ul", 142, 849, 278, 877),
        ]
    ensure_dir(output_folder)

    for filename in os.listdir(input_folder):
        if not is_image(filename):
            continue
        image = _cv2_imread(cv2, os.path.join(input_folder, filename))
        if image is None:
            continue

        h, w, _ = image.shape
        base_name = os.path.splitext(filename)[0]

        for name, x1, y1, x2, y2 in regions:
            x2c, y2c = min(x2, w), min(y2, h)
            x1c, y1c = max(x1, 0), max(y1, 0)
            cropped = image[y1c:y2c, x1c:x2c]
            if cropped.size > 0:
                out_path = os.path.join(output_folder, f"{base_name}_{name}.jpg")
                if not _cv2_imwrite(cv2, out_path, cropped):
                    print(f"警告: 保存图片失败 {out_path}")

    print("第二步：主数据截图裁剪完成")


# ======================== 第三步：OCR 识别与数据提取 ========================

def step3_ocr_and_extract(image_folder, result_excel, finish_excel,
                          target_freq_prefix=525150, prefix_length=6,
                          concurrency=None):
    """OCR 识别图片，提取数据到 Excel
    concurrency: 并发引擎数，None 时使用全局 _OCR_CONCURRENCY"""
    logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")

    workers = _OCR_CONCURRENCY if concurrency is None else concurrency
    if workers < 1:
        workers = 1

    # ---- 3a. OCR 识别 ----
    image_files = [f for f in os.listdir(image_folder) if is_image(f)]
    if not image_files:
        logging.warning("未找到符合格式的图片")
        return

    results = []
    with ThreadPoolExecutor(max_workers=workers) as executor:
        future_map = {
            executor.submit(ocr_image, os.path.join(image_folder, f),
                            thread_engine=True): f
            for f in image_files
        }
        for future in as_completed(future_map):
            img_name = future_map[future]
            try:
                text = future.result()
                formatted = format_ocr_text(text)
                results.append([img_name, formatted])
                logging.info(f"处理完成: {img_name}")
            except Exception as e:
                logging.error(f"处理 {img_name} 失败: {e}")

    results.sort(key=lambda x: x[0])
    _write_excel_sheet(result_excel, "Sheet1", results, ["文件名", "识别文本"],
                       clear_data=True)

    # ---- 3b. 数据提取 ----
    extract_result_to_finish(result_excel, finish_excel,
                             target_freq_prefix, prefix_length)
    print(f"第三步完成：结果保存到 {finish_excel}")


def format_ocr_text(text):
    """格式化 OCR 文本，数字保留两位小数（不四舍五入）"""
    try:
        words = text.split()
        result = []
        for word in words:
            try:
                result.append(str(Decimal(word).quantize(Decimal('0.00'), rounding=ROUND_DOWN)))
            except (ValueError, ArithmeticError):
                result.append(word)
        return " ".join(result)
    except Exception:
        return text


def _write_excel_sheet(excel_path, sheet_name, data_rows, columns,
                       clear_data=False):
    """将数据写入 Excel 工作表（自动处理新文件/已有文件）"""
    if clear_data and os.path.exists(excel_path):
        wb = load_workbook(excel_path)
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row in ws.iter_rows(min_row=2, max_col=ws.max_column, max_row=ws.max_row):
                for cell in row:
                    cell.value = None
            wb.save(excel_path)
        wb.close()

    mode = 'a' if os.path.exists(excel_path) else 'w'
    with pd.ExcelWriter(excel_path, engine='openpyxl', mode=mode,
                        if_sheet_exists='replace') as writer:
        pd.DataFrame(data_rows, columns=columns).to_excel(
            writer, sheet_name=sheet_name, index=False)


def _build_files_info(df):
    """从 DataFrame 构建 {filename: {value, filepath}} 字典"""
    # 确保至少3列（兼容只有2列的 OCR 结果文件）
    while df.shape[1] < 3:
        df = pd.concat([df, pd.DataFrame(columns=[df.shape[1]])], axis=1)

    info = {}
    for _, row in df.iterrows():
        if pd.notna(row[0]):
            info[str(row[0]).strip()] = {
                'value': row[1] if pd.notna(row[1]) else '',
                'filepath': row[2] if pd.notna(row[2]) else '',
            }
    return info


def extract_ping_data(file_path):
    """提取 ping 延迟和抖动数据"""
    if not os.path.exists(file_path):
        print(f"错误：找不到文件 {file_path}")
        return []
    try:
        df = pd.read_excel(file_path, header=None)
        files_info = _build_files_info(df)

        categories = {
            'ping32_delay': (r'^(\d+)ping32_delay\.jpg$', {}),
            'ping32_jitter': (r'^(\d+)ping32_jitter\.jpg$', {}),
            'ping2000_delay': (r'^(\d+)ping2000_delay\.jpg$', {}),
            'ping2000_jitter': (r'^(\d+)ping2000_jitter\.jpg$', {}),
        }

        for filename, info in files_info.items():
            for _cat_name, (pattern, storage) in categories.items():
                m = re.search(pattern, filename)
                if m:
                    val = info['value']
                    if val != '' and pd.notna(val):
                        try:
                            v = float(val)
                            storage[m.group(1)] = v if v == int(v) else float(f"{v:.2f}")
                        except ValueError:
                            print(f"值转换错误: {filename} -> {val}")
                    break

        all_groups = set().union(*[d.keys() for _, (_, d) in categories.items()])
        return [
            {
                'Group': g,
                'ping32_delay': categories['ping32_delay'][1].get(g),
                'ping32_jitter': categories['ping32_jitter'][1].get(g),
                'ping2000_delay': categories['ping2000_delay'][1].get(g),
                'ping2000_jitter': categories['ping2000_jitter'][1].get(g),
            }
            for g in sorted(all_groups,
                            key=lambda x: (0, int(x)) if x.isdigit() else (1, x))
        ]
    except Exception as e:
        import traceback
        print(f"读取文件时发生错误: {e}")
        traceback.print_exc()
        return []


def _safe_float(val):
    """安全转换为 float，失败返回 None"""
    try:
        return float(val)
    except (ValueError, TypeError):
        return None


def extract_main_data(file_path, target_freq_prefix, prefix_length=6):
    """提取主数据（频点、RSRP、SINR、DL、UL）"""
    if not os.path.exists(file_path):
        print(f"错误：找不到文件 {file_path}")
        return []
    try:
        df = pd.read_excel(file_path, header=None)
        files_info = _build_files_info(df)

        data_map = {}
        for filename, info in files_info.items():
            val = info['value']
            if val == '' or pd.isna(val):
                continue

            num_val = _safe_float(val)
            if num_val is None:
                continue

            for suffix, key in [('_pindian_', 'pindian'), ('_rsrp_', 'rsrp'),
                                ('_sinr_', 'sinr')]:
                m = re.search(rf'(\w+){suffix}(\d+)\.jpg', filename)
                if m:
                    data_map.setdefault(key, {})[f"{m.group(1)}_{m.group(2)}"] = num_val
                    break

            for suffix, key in [('_dl', 'dl'), ('_ul', 'ul')]:
                m = re.search(rf'(\w+){suffix}\.jpg', filename)
                if m:
                    data_map.setdefault(key, {})[m.group(1)] = num_val

        pindian = data_map.get('pindian', {})
        rsrp = data_map.get('rsrp', {})
        sinr = data_map.get('sinr', {})
        dl = data_map.get('dl', {})
        ul = data_map.get('ul', {})

        valid_configs = set(pindian) | set(rsrp) | set(sinr)
        results = []
        for config in sorted(valid_configs):
            freq = pindian.get(config)
            if freq is None:
                continue
            freq_int = int(freq)
            freq_str = str(freq_int)
            freq_prefix = int(freq_str[:prefix_length]) if len(freq_str) > prefix_length else freq_int
            if freq_prefix != target_freq_prefix:
                continue

            prefix = '_'.join(config.split('_')[:-1])
            results.append({
                '配置': config,
                '频点': freq_prefix,
                'RSRP': _truncate(rsrp.get(config), 2),
                'SINR': _truncate(sinr.get(config), 2),
                'DL': _truncate(dl.get(prefix), 2),
                'UL': _truncate(ul.get(prefix), 2),
            })
        return results
    except Exception as e:
        import traceback
        print(f"读取文件时发生错误: {e}")
        traceback.print_exc()
        return []


def extract_rank_data(file_path):
    """从 result.xlsx 提取 rank 值，返回 {prefix: rank_value}"""
    if not os.path.exists(file_path):
        print(f"错误：找不到文件 {file_path}")
        return {}
    try:
        df = pd.read_excel(file_path, header=None)
        rank_dict = {}
        for _, row in df.iterrows():
            fname = str(row[0])
            if '_rank' in fname:
                prefix = fname.split('_rank')[0]
                try:
                    rank_dict[prefix] = float(row[1])
                except (ValueError, TypeError):
                    pass
        print(f"提取到 {len(rank_dict)} 条 Rank 数据")
        return rank_dict
    except Exception as e:
        print(f"提取 Rank 数据失败: {e}")
        return {}


def extract_result_to_finish(result_excel, finish_excel,
                             target_freq_prefix=525150, prefix_length=6):
    """从已有的 result.xlsx 提取数据并写入 finish.xlsx，不执行 OCR"""
    if not os.path.exists(result_excel):
        raise FileNotFoundError(f"找不到 result.xlsx: {result_excel}")

    main_data = extract_main_data(result_excel, target_freq_prefix, prefix_length)
    ping_data = extract_ping_data(result_excel)
    rank_dict = extract_rank_data(result_excel)

    for row in main_data:
        config = row.get('配置', '')
        prefix = config.rsplit('_', 1)[0] if '_' in config else config
        row['Rank'] = rank_dict.get(prefix)

    out_dir = os.path.dirname(finish_excel)
    if out_dir:
        ensure_dir(out_dir)

    with pd.ExcelWriter(finish_excel, engine='openpyxl') as writer:
        pd.DataFrame(main_data).to_excel(writer, sheet_name='Sheet1', index=False)
        pd.DataFrame(ping_data).to_excel(writer, sheet_name='Sheet2', index=False)

    rank_matched = sum(1 for r in main_data if r.get('Rank') is not None)
    print(f"提取完成：{finish_excel}")
    print(f"  Sheet1 频点数据: {len(main_data)} 行 (Rank: {rank_matched} 条匹配)")
    print(f"  Sheet2 ping 数据: {len(ping_data)} 行")
    return main_data, ping_data


def _truncate(value, decimals):
    """截断到指定小数位，不四舍五入"""
    if value is None:
        return None
    s = f"{value:.10f}"
    if '.' in s:
        integer, decimal = s.split('.')
        return float(f"{integer}.{decimal[:decimals]}") if decimals > 0 else float(integer)
    return float(s)


# ======================== 第四步：插入主数据截图 ========================

# 每组图片的参数：起始行、hd宽度、hd高度、其他图宽度、其他图高度
_MAIN_SCREENSHOT_CONFIGS = [
    # (row, hd_w, hd_h, other_w, other_h)  — 单位 cm，96 DPI
    (6,   10.2, 12.8, 8.7, 12.8),
    (28,  10.2, 12.6, 8.7, 12.6),
    (50,  10.2, 12.8, 8.7, 12.8),
    (72,  10.2, 12.6, 8.7, 12.6),
    (94,  10.2, 12.8, 8.7, 12.8),
    (116, 10.2, 12.6, 8.7, 12.6),
    (138, 10.2, 12.8, 8.7, 12.8),
    (160, 10.2, 12.6, 8.7, 12.6),
    (182, 10.2, 12.8, 8.7, 12.8),
    (204, 10.2, 12.6, 8.7, 12.6),
]

def step4_insert_main_screenshots(template_excel, num_groups=6, configs=None,
                                   img_columns=None, pictuer_dir=None,
                                   sheet_name='SA-网络性能验收-CQT',
                                   positions_override=None):
    """将 pictuer 截图插入到模板的指定工作表
    num_groups: 插入组数
    configs: [(row, hd_w, hd_h, other_w, other_h), ...]  默认使用预定义配置
    img_columns: [(suffix, col), ...]  默认 [('hd','AC'),('hu','AT'),('cd','BK'),('cu','CB')]
    pictuer_dir: pictuer 图片目录，默认使用全局 PICTUER_DIR
    sheet_name: 目标工作表名
    positions_override: 直接传入位置字典，跳过参数计算"""
    if positions_override is not None:
        positions = positions_override
    else:
        if pictuer_dir is None:
            pictuer_dir = PICTUER_DIR
        if configs is None:
            configs = _MAIN_SCREENSHOT_CONFIGS
        if img_columns is None:
            img_columns = [('hd', 'AC'), ('hu', 'AT'), ('cd', 'BK'), ('cu', 'CB')]

        positions = {}
        for i, (row, hd_w, hd_h, o_w, o_h) in enumerate(configs[:num_groups]):
            g = str(i + 1)
            for suffix, col in img_columns:
                w, h = (hd_w, hd_h) if suffix == 'hd' else (o_w, o_h)
                path = os.path.join(pictuer_dir, f"{g}{suffix}.png")
                if os.path.exists(path):
                    positions[path] = {"position": f"{col}{row}", "width_cm": w, "height_cm": h}

    wb = _open_template(template_excel)
    insert_images_to_excel(wb, sheet_name, positions)
    wb.save(template_excel)
    print(f"第四步：主数据截图已插入 {len(positions)} 张到 {template_excel}")


# ======================== 第五步：插入 ping 截图 ========================

def step5_insert_ping_screenshots(template_excel, num_groups=6, start_row=6,
                                   row_step=22, img_columns=None,
                                   width_cm=8.6, height_cm=8.5, ping_dir=None,
                                   sheet_name='SA-网络性能验收-CQT',
                                   positions_override=None):
    """将 ping 截图插入到模板的指定工作表
    num_groups: 插入组数
    img_columns: [(suffix, col), ...]  默认 [('ping32','CS'),('ping2000','DJ')]
    ping_dir: ping 图片目录，默认使用全局 PING_DIR
    sheet_name: 目标工作表名
    positions_override: 直接传入位置字典，跳过参数计算"""
    if positions_override is not None:
        positions = positions_override
    else:
        if ping_dir is None:
            ping_dir = PING_DIR
        if img_columns is None:
            img_columns = [('ping32', 'CS'), ('ping2000', 'DJ')]

        positions = {}
        for i in range(1, num_groups + 1):
            row = start_row + (i - 1) * row_step
            for suffix, col in img_columns:
                path = os.path.join(ping_dir, f"{i}{suffix}.png")
                if os.path.exists(path):
                    positions[path] = {"position": f"{col}{row}", "width_cm": width_cm, "height_cm": height_cm}

    wb = _open_template(template_excel)
    insert_images_to_excel(wb, sheet_name, positions, clear_first=False)
    wb.save(template_excel)
    print(f"第五步：ping 截图已插入 {len(positions)} 张到 {template_excel}")


# ======================== 第六步：插入遍历测试图 ========================

def step6_insert_dt_screenshots(template_excel, start_num=1, end_num=6,
                                img_types=None, img_cols=None,
                                start_row=7, row_step=18,
                                width_cm=10.2, height_cm=7.79, dt_dir=None,
                                sheet_name='SA-网络性能验收-遍历测试图',
                                positions_override=None):
    """将 DT 遍历测试图插入到模板的指定工作表
    img_types/img_cols: 图片类型与对应列
    dt_dir: DT 图片目录，默认使用全局 DT_DIR
    sheet_name: 目标工作表名
    positions_override: 直接传入位置字典，跳过参数计算"""
    if positions_override is not None:
        positions = positions_override
    else:
        if dt_dir is None:
            dt_dir = DT_DIR
        if img_types is None:
            img_types = ["RSRP", "SINR", "PCI", "DL", "UL"]
        if img_cols is None:
            img_cols = ["C", "I", "AA", "O", "U"]

        positions = {}
        for idx, folder_name in enumerate(range(start_num, end_num + 1)):
            folder_path = os.path.join(dt_dir, str(folder_name))
            if not os.path.exists(folder_path):
                print(f"警告: 文件夹 {folder_path} 不存在，跳过。")
                continue
            row = start_row + idx * row_step
            for img_type, col in zip(img_types, img_cols):
                path = os.path.join(folder_path, f"{img_type}.png")
                if os.path.exists(path):
                    positions[path] = {"position": f"{col}{row}", "width_cm": width_cm, "height_cm": height_cm}

    wb = _open_template(template_excel)
    insert_images_to_excel(wb, sheet_name, positions)
    wb.save(template_excel)
    print(f"第六步：遍历测试图已插入 {len(positions)} 张到 {template_excel}")

    clear_folder(JIETU_DIR)
    print("已清空 jietu 文件夹")


# ======================== 可编辑 Treeview ========================

class EditableTreeview(ttk.Treeview):
    """支持双击编辑单元格的 Treeview"""

    def __init__(self, parent, columns, col_names, col_widths=None, **kw):
        super().__init__(parent, columns=columns, show="headings", **kw)
        self.col_keys = columns
        for i, name in enumerate(col_names):
            w = col_widths[i] if col_widths else 100
            self.heading(columns[i], text=name)
            self.column(columns[i], width=w, anchor="center")
        self.tag_configure("even", background="#FFFFFF")
        self.tag_configure("odd", background="#F2F7FA")
        self._edit_cell = None
        self.bind("<Double-1>", self._on_double_click)

    def _on_double_click(self, event):
        region = self.identify_region(event.x, event.y)
        if region != "cell":
            return
        col = self.identify_column(event.x)
        item = self.identify_row(event.y)
        if not item:
            return
        col_idx = int(col[1:]) - 1
        x, y, w, h = self.bbox(item, col)
        old_val = self.set(item, self.col_keys[col_idx])
        entry = ttk.Entry(self)
        entry.place(x=x, y=y, width=w, height=h)
        entry.insert(0, old_val)
        entry.select_range(0, "end")
        entry.focus_set()

        def save(event=None):
            new_val = entry.get()
            self.set(item, self.col_keys[col_idx], new_val)
            entry.destroy()
            self._edit_cell = None

        entry.bind("<Return>", save)
        entry.bind("<FocusOut>", save)
        entry.bind("<Escape>", lambda e: entry.destroy())
        self._edit_cell = entry

    def add_row(self, values, tags=()):
        return self.insert("", "end", values=values, tags=tags)

    def delete_selected(self):
        for item in self.selection():
            self.delete(item)

    def get_all_rows(self):
        return [[self.set(item, k) for k in self.col_keys]
                for item in self.get_children()]

    def clear(self):
        for item in self.get_children():
            self.delete(item)

    def load_rows(self, rows):
        self.clear()
        for i, row in enumerate(rows):
            self.add_row(row, tags=("odd",) if i % 2 else ("even",))


# ======================== 主界面 ========================

class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("报告输出")
        self.minsize(850, 620)
        self._center_window(980, 750)

        # 简单变量
        self.base_dir = tk.StringVar(value=r"D:\baogao\sferzaibo")
        self.template_path = tk.StringVar(value="")
        self.cqt_sheet_name = tk.StringVar(value="SA-网络性能验收-CQT")
        self.dt_sheet_name = tk.StringVar(value="SA-网络性能验收-遍历测试图")
        self.freq_prefix = tk.StringVar(value="525150")
        self.freq_length = tk.IntVar(value=6)
        self.ocr_concurrency = tk.IntVar(value=_OCR_CONCURRENCY)
        self.main_groups = tk.IntVar(value=4)
        self.dt_start = tk.IntVar(value=1)
        self.dt_end = tk.IntVar(value=6)
        self.profile_name = tk.StringVar(value="默认")

        # 复杂变量（在 _build_* 中创建）
        self.ping_start_row = tk.StringVar(value="6")
        self.ping_row_step = tk.StringVar(value="22")
        self.ping_w = tk.StringVar(value="8.6")
        self.ping_h = tk.StringVar(value="8.5")
        self.ping_cols = tk.StringVar(value="ping32,CS ping2000,DJ")
        self.dt_start_row = tk.StringVar(value="7")
        self.dt_row_step = tk.StringVar(value="18")
        self.dt_w = tk.StringVar(value="10.2")
        self.dt_h = tk.StringVar(value="7.79")
        self.dt_cols = tk.StringVar(value="RSRP,C SINR,I PCI,AA DL,O UL,U")
        self.cqt_main_cols = tk.StringVar(value="hd,AC hu,AT cd,BK cu,CB")
        self._ui_queue = queue.Queue()
        self._running = False
        self._run_error = False
        self._cancel = threading.Event()
        self._error_shown = False

        self._setup_style()
        self._make_app_icon()
        self._build_ui()
        self._load_profiles()

        self.after(100, self._poll_ui_queue)
        self.protocol("WM_DELETE_WINDOW", self._on_close)

    # ==================== UI 构建 ====================

    def _center_window(self, w, h):
        """让窗口在屏幕居中显示。"""
        sw = self.winfo_screenwidth()
        sh = self.winfo_screenheight()
        x = max(0, (sw - w) // 2)
        y = max(0, (sh - h) // 2)
        self.geometry(f"{w}x{h}+{x}+{y}")

    def _setup_style(self):
        ui_font = "Microsoft YaHei UI"
        self._colors = {
            "bg": "#F1F5F9",           # 页面背景（浅灰蓝）
            "card": "#FFFFFF",         # 卡片背景
            "border": "#E2E8F0",       # 边框
            "header": "#1E3A8A",       # 标题栏（深蓝）
            "header_text": "#FFFFFF",  # 标题文字
            "subtitle": "#BFDBFE",     # 副标题（浅蓝）
            "footer": "#F8FAFC",       # 底部状态栏
            "text": "#1E293B",         # 正文文字
            "muted": "#64748B",        # 次要文字
            "heading": "#334155",      # 标题文字
            "heading_bg": "#E2E8F0",   # 表头背景
            "accent": "#2563EB",       # 强调色（亮蓝）
            "accent_dark": "#1D4ED8",  # 强调色深
            "accent_light": "#DBEAFE", # 强调色浅
            "button": "#FFFFFF",
            "button_active": "#EFF6FF",
            "button_disabled": "#E2E8F0",
            "tab": "#E2E8F0",
            "tab_active": "#EFF6FF",
            "trough": "#E2E8F0",
            "success": "#16A34A",
            "running": "#F59E0B",
            "error": "#DC2626",
        }
        self._fonts = {
            "title": (ui_font, 12, "bold"),
            "subtitle": (ui_font, 8),
            "body": (ui_font, 9),
            "tab": (ui_font, 9),
            "section": (ui_font, 9, "bold"),
            "heading": (ui_font, 8, "bold"),
            "button_bold": (ui_font, 9, "bold"),
            "mono": ("Consolas", 9),
        }

        style = ttk.Style(self)
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass

        c = self._colors
        f = self._fonts
        self.configure(bg=c["bg"])

        style.configure(".", font=f["body"], foreground=c["text"])
        style.configure("TFrame", background=c["bg"])
        style.configure("Card.TFrame", background=c["card"])
        style.configure("Header.TFrame", background=c["header"])
        style.configure("Footer.TFrame", background=c["footer"])

        style.configure("Title.TLabel", background=c["header"],
                        foreground=c["header_text"], font=f["title"])
        style.configure("Subtitle.TLabel", background=c["header"],
                        foreground=c["subtitle"], font=f["subtitle"])
        style.configure("TLabel", background=c["bg"], foreground=c["text"])
        style.configure("Card.TLabel", background=c["card"], foreground=c["text"])
        style.configure("Hint.TLabel", background=c["card"], foreground=c["muted"])
        style.configure("Status.TLabel", background=c["footer"],
                        foreground=c["text"], font=f["mono"])

        style.configure("TButton", padding=(10, 5), background=c["button"],
                        foreground=c["text"], borderwidth=1, focusthickness=0)
        style.map("TButton",
                  background=[("active", c["button_active"]),
                              ("disabled", c["button_disabled"])],
                  foreground=[("disabled", c["muted"])])
        style.configure("Accent.TButton", background=c["accent"],
                        foreground="#FFFFFF", font=f["button_bold"])
        style.map("Accent.TButton",
                  background=[("active", c["accent_dark"]),
                              ("disabled", c["button_disabled"])],
                  foreground=[("disabled", "#C9D6DB")])

        style.configure("TNotebook", background=c["bg"], borderwidth=0)
        style.configure("TNotebook.Tab", padding=(12, 5), background=c["tab"],
                        foreground=c["text"], borderwidth=0, font=f["tab"])
        style.map("TNotebook.Tab",
                  background=[("selected", c["card"]),
                              ("active", c["tab_active"])],
                  foreground=[("selected", c["accent_dark"])])

        style.configure("TLabelframe", background=c["card"],
                        bordercolor=c["border"], relief="solid", borderwidth=1)
        style.configure("TLabelframe.Label", background=c["card"],
                        foreground=c["heading"], font=f["section"])

        style.configure("TEntry", fieldbackground="#FFFFFF",
                        foreground=c["text"], bordercolor=c["border"],
                        lightcolor=c["border"], darkcolor=c["border"], padding=3)
        style.map("TEntry", bordercolor=[("focus", c["accent"])])
        style.configure("TCombobox", fieldbackground="#FFFFFF",
                        foreground=c["text"], background=c["button"],
                        bordercolor=c["border"], lightcolor=c["border"],
                        darkcolor=c["border"], padding=3)
        style.map("TCombobox", bordercolor=[("focus", c["accent"])])
        style.configure("TSpinbox", fieldbackground="#FFFFFF",
                        foreground=c["text"], bordercolor=c["border"],
                        lightcolor=c["border"], darkcolor=c["border"], padding=3)

        style.configure("Treeview", background="#FFFFFF",
                        fieldbackground="#FFFFFF", foreground=c["text"],
                        rowheight=23, bordercolor=c["border"])
        style.map("Treeview",
                  background=[("selected", c["accent_light"])],
                  foreground=[("selected", c["text"])])
        style.configure("Treeview.Heading", background=c["heading_bg"],
                        foreground=c["heading"], font=f["heading"],
                        padding=(6, 4), relief="flat")

        style.configure("TProgressbar", troughcolor=c["trough"],
                        background=c["accent"], bordercolor=c["trough"],
                        lightcolor=c["accent"], darkcolor=c["accent"],
                        thickness=12)

    def _make_app_icon(self):
        img = PilImage.new("RGBA", (64, 64), (23, 59, 92, 255))
        draw = ImageDraw.Draw(img)
        draw.rectangle([4, 4, 59, 59], fill=(23, 59, 92, 255),
                       outline=(255, 255, 255, 60), width=2)
        for x, top in [(14, 36), (23, 30), (32, 24), (41, 18)]:
            draw.rectangle([x, top, x + 5, 46], fill=(72, 196, 198, 255))
        draw.ellipse([49, 14, 57, 22], fill=(255, 199, 90, 255))
        self._app_icon = ImageTk.PhotoImage(img)
        try:
            self.iconphoto(True, self._app_icon)
        except tk.TclError:
            pass

    def _build_ui(self):
        # ---- 顶部标题栏 ----
        header = ttk.Frame(self, style="Header.TFrame")
        header.pack(fill="x")
        accent_line = tk.Frame(self, bg=self._colors["accent"], height=2)
        accent_line.pack(fill="x")
        title_box = ttk.Frame(header, style="Header.TFrame")
        title_box.pack(side="left", fill="x", expand=True, padx=12, pady=5)
        ttk.Label(title_box, text="报告输出", style="Title.TLabel").pack(side="left")
        ttk.Label(title_box, text="5G SA 网络性能验收报告生成",
                  style="Subtitle.TLabel").pack(side="left", padx=(10, 0))
        self.start_btn = ttk.Button(header, text="▶  开始执行",
                                    style="Accent.TButton", command=self._start_run)
        self.start_btn.pack(side="right", padx=12)
        self.stop_btn = ttk.Button(header, text="■  停止",
                                   style="Accent.TButton", command=self._stop_run)
        self.stop_btn.pack(side="right", padx=(0, 4))
        self.stop_btn.config(state="disabled")

        # ---- 配置方案 + 工作目录 ----
        top = ttk.Frame(self, style="Card.TFrame", padding=10)
        top.pack(fill="x", padx=10, pady=(10, 0))

        profile_bar = ttk.Frame(top, style="Card.TFrame")
        profile_bar.pack(fill="x")
        ttk.Label(profile_bar, text="配置方案:", style="Card.TLabel").pack(side="left")
        self.profile_combo = ttk.Combobox(profile_bar, textvariable=self.profile_name,
                                          width=16, state="readonly")
        self.profile_combo.pack(side="left", padx=6)
        self.profile_combo.bind("<<ComboboxSelected>>", self._on_profile_select)
        ttk.Button(profile_bar, text="保存", command=self._save_profile).pack(side="left", padx=2)
        ttk.Button(profile_bar, text="另存为...",
                   command=self._save_as_profile).pack(side="left")
        ttk.Button(profile_bar, text="删除",
                   command=self._delete_profile).pack(side="left", padx=2)

        dir_bar = ttk.Frame(top, style="Card.TFrame")
        dir_bar.pack(fill="x", pady=(6, 0))
        ttk.Label(dir_bar, text="工作目录:", style="Card.TLabel").pack(side="left")
        ttk.Entry(dir_bar, textvariable=self.base_dir, width=56).pack(side="left", padx=6)
        ttk.Button(dir_bar, text="浏览...", command=self._browse_dir).pack(side="left")
        ttk.Label(dir_bar, text="子目录将自动派生",
                  style="Hint.TLabel").pack(side="left", padx=10)

        # ---- 底部：状态与进度 ----
        bottom = ttk.Frame(self, style="Footer.TFrame", padding=(10, 7))
        bottom.pack(side="bottom", fill="x", padx=10, pady=(0, 10))

        status_row = ttk.Frame(bottom, style="Footer.TFrame")
        status_row.pack(fill="x")
        self.status_dot = tk.Canvas(status_row, width=12, height=12,
                                    bg=self._colors["footer"], highlightthickness=0)
        self.status_dot.pack(side="left", padx=(0, 6))
        self._status_dot_id = self.status_dot.create_oval(
            1, 1, 11, 11, fill=self._colors["muted"], outline="")
        self.status_var = tk.StringVar(value="就绪")
        ttk.Label(status_row, textvariable=self.status_var,
                  style="Status.TLabel").pack(side="left", padx=(0, 10))

        self.progress = ttk.Progressbar(bottom, mode="determinate", maximum=100)
        self.progress.pack(fill="x", pady=(6, 0))

        # ---- 标签页 ----
        nb = ttk.Notebook(self)
        nb.pack(fill="both", expand=True, padx=10, pady=(10, 0))
        self._build_crop_tab(nb)
        self._build_ocr_tab(nb)
        self._build_image_tab(nb)
        self._build_params_tab(nb)

    def _set_status_dot(self, color):
        if hasattr(self, "status_dot") and self._status_dot_id:
            try:
                self.status_dot.itemconfig(self._status_dot_id, fill=color)
            except tk.TclError:
                pass

    # ---- 标签页构建 ----

    def _build_crop_tab(self, nb):
        tab = ttk.Frame(nb)
        nb.add(tab, text="数据位置")

        crop_nb = ttk.Notebook(tab)
        crop_nb.pack(fill="both", expand=True)

        # 初始化两个预览器的状态
        self._pv = {"ping": {}, "main": {}}

        # ---- Ping 裁剪 ----
        ping_frame = ttk.Frame(crop_nb)
        crop_nb.add(ping_frame, text="Ping 范围")
        self._build_crop_section(ping_frame, "ping",
            default_rows=[("delay","572","30","639","51"),
                          ("jitter","506","79","567","95")])

        # ---- 主数据裁剪 ----
        main_frame = ttk.Frame(crop_nb)
        crop_nb.add(main_frame, text="主数据范围")
        self._build_crop_section(main_frame, "main",
            default_rows=[
                ("pindian_1","137","229","198","251"),
                ("rsrp_1","449","29","512","53"),
                ("sinr_1","451","55","513","77"),
                ("pindian_2","226","229","283","251"),
                ("rsrp_2","516","29","578","52"),
                ("sinr_2","516","55","579","77"),
                ("dl","141","648","276","675"),
                ("ul","142","849","278","877")])

    def _build_crop_section(self, parent, key, default_rows):
        """构建一个裁剪区域（左表格 + 右图片预览）"""
        paned = ttk.PanedWindow(parent, orient="horizontal")
        paned.pack(fill="both", expand=True)

        # ---- 左侧：表格 ----
        left = ttk.Frame(paned)
        paned.add(left, weight=1)

        tree = EditableTreeview(
            left, columns=["name","x1","y1","x2","y2"],
            col_names=["区域名称","x1","y1","x2","y2"],
            col_widths=[100,65,65,65,65])
        tree.pack(fill="both", expand=True, padx=2, pady=2)
        tree.load_rows(default_rows)
        setattr(self, f"_{key}_tree", tree)

        bar = ttk.Frame(left)
        bar.pack(fill="x", padx=2, pady=2)
        ttk.Button(bar, text="+",
                   command=lambda: tree.add_row(["","0","0","0","0"])).pack(side="left")
        ttk.Button(bar, text="−",
                   command=tree.delete_selected).pack(side="left", padx=2)
        ttk.Button(bar, text="填入选中行",
                   command=lambda: self._fill_preview_to_tree(key, tree)).pack(side="left", padx=4)

        # ---- 右侧：图片预览 ----
        right = ttk.Frame(paned)
        paned.add(right, weight=1)

        load_bar = ttk.Frame(right)
        load_bar.pack(fill="x", padx=2, pady=2)
        ttk.Button(load_bar, text="加载图片...",
                   command=lambda: self._load_crop_image(key)).pack(side="left")
        ttk.Button(load_bar, text="适合窗口",
                   command=lambda: self._fit_crop_view(key)).pack(side="left", padx=4)

        # Canvas
        cf = ttk.Frame(right)
        cf.pack(fill="both", expand=True, padx=2, pady=2)
        canvas = tk.Canvas(cf, bg="#2d2d2d", cursor="crosshair")
        hs = ttk.Scrollbar(cf, orient="horizontal", command=canvas.xview)
        vs = ttk.Scrollbar(cf, orient="vertical", command=canvas.yview)
        canvas.configure(xscrollcommand=hs.set, yscrollcommand=vs.set)
        canvas.grid(row=0, column=0, sticky="nsew")
        hs.grid(row=1, column=0, sticky="ew")
        vs.grid(row=0, column=1, sticky="ns")
        cf.rowconfigure(0, weight=1)
        cf.columnconfigure(0, weight=1)

        # 信息标签
        info_frame = ttk.Frame(right)
        info_frame.pack(fill="x", padx=2, pady=2)
        coord_label = tk.StringVar(value="鼠标: —")
        ttk.Label(info_frame, textvariable=coord_label, font=("Consolas", 10),
                  foreground="blue").pack(side="left")
        point1_label = tk.StringVar(value="点1: —")
        ttk.Label(info_frame, textvariable=point1_label, font=("Consolas", 10),
                  foreground="green").pack(side="left", padx=12)
        point2_label = tk.StringVar(value="点2: —")
        ttk.Label(info_frame, textvariable=point2_label, font=("Consolas", 10),
                  foreground="red").pack(side="left", padx=12)

        # 绑定事件
        canvas.bind("<Motion>",
                    lambda e, k=key: self._on_crop_motion(k, e))
        canvas.bind("<Button-1>",
                    lambda e, k=key: self._on_crop_click(k, e))
        canvas.bind("<MouseWheel>",
                    lambda e, k=key: self._on_crop_wheel(k, e))

        # 存储状态
        self._pv[key] = {
            "canvas": canvas,
            "image": None,      # PIL Image
            "tk_img": None,     # PhotoImage
            "scale": 1.0,
            "click1": None,     # (x, y) image coords
            "click2": None,
            "coord_label": coord_label,
            "point1_label": point1_label,
            "point2_label": point2_label,
        }

    # ========== 图片预览（两点点击模式）==========

    def _load_crop_image(self, key):
        path = filedialog.askopenfilename(
            title="选择截图",
            filetypes=[("图片", "*.png *.jpg *.jpeg *.bmp"), ("所有文件", "*.*")])
        if not path:
            return
        pv = self._pv[key]
        pv["image"] = PilImage.open(path)
        pv["scale"] = 1.0
        pv["click1"] = None
        pv["click2"] = None
        pv["point1_label"].set("点1: —")
        pv["point2_label"].set("点2: —")
        self._redraw_crop(key)

    def _fit_crop_view(self, key):
        pv = self._pv[key]
        if not pv["image"]:
            return
        cw = pv["canvas"].winfo_width()
        ch = pv["canvas"].winfo_height()
        iw, ih = pv["image"].size
        if cw > 10 and ch > 10:
            pv["scale"] = min(cw / iw, ch / ih, 1.0)
        self._redraw_crop(key)

    def _redraw_crop(self, key):
        pv = self._pv[key]
        if not pv["image"]:
            return
        w = int(pv["image"].width * pv["scale"])
        h = int(pv["image"].height * pv["scale"])
        img = pv["image"].resize((w, h), PilImage.LANCZOS)
        pv["tk_img"] = ImageTk.PhotoImage(img)
        pv["canvas"].delete("all")
        pv["canvas"].create_image(0, 0, anchor="nw", image=pv["tk_img"])
        pv["canvas"].configure(scrollregion=(0, 0, w, h))

        # 重绘已有的点和选框
        for i in [1, 2]:
            pt = pv.get(f"click{i}")
            if pt:
                cx = pt[0] * pv["scale"]
                cy = pt[1] * pv["scale"]
                color = "green" if i == 1 else "red"
                r = 3
                pv["canvas"].create_oval(cx - r, cy - r, cx + r, cy + r,
                                         outline=color, width=2, tags=f"pt{i}")
        # 选框
        if pv["click1"] and pv["click2"]:
            x1 = pv["click1"][0] * pv["scale"]
            y1 = pv["click1"][1] * pv["scale"]
            x2 = pv["click2"][0] * pv["scale"]
            y2 = pv["click2"][1] * pv["scale"]
            pv["canvas"].create_rectangle(x1, y1, x2, y2, outline="yellow",
                                          width=2, dash=(4, 2), tags="rect")

    def _canvas_to_img(self, key, cx, cy):
        pv = self._pv[key]
        if not pv["image"]:
            return None, None
        x = int(pv["canvas"].canvasx(cx) / pv["scale"])
        y = int(pv["canvas"].canvasy(cy) / pv["scale"])
        return x, y

    def _on_crop_motion(self, key, event):
        x, y = self._canvas_to_img(key, event.x, event.y)
        if x is not None:
            self._pv[key]["coord_label"].set(f"鼠标: ({x}, {y})")

    def _on_crop_click(self, key, event):
        pv = self._pv[key]
        if not pv["image"]:
            return
        x, y = self._canvas_to_img(key, event.x, event.y)
        if x is None:
            return

        if pv["click1"] is None:
            # 第一点
            pv["click1"] = (x, y)
            pv["click2"] = None
            pv["point1_label"].set(f"点1: ({x}, {y})")
            pv["point2_label"].set("点2: —")
        elif pv["click2"] is None:
            # 第二点
            # 确保 x1<=x2, y1<=y2
            x1, y1 = pv["click1"]
            pv["click1"] = (min(x1, x), min(y1, y))
            pv["click2"] = (max(x1, x), max(y1, y))
            pv["point1_label"].set(f"点1: ({pv['click1'][0]}, {pv['click1'][1]})")
            pv["point2_label"].set(f"点2: ({pv['click2'][0]}, {pv['click2'][1]})")
        else:
            # 重置，重新开始
            pv["click1"] = (x, y)
            pv["click2"] = None
            pv["point1_label"].set(f"点1: ({x}, {y})")
            pv["point2_label"].set("点2: —")

        self._redraw_crop(key)

    def _on_crop_wheel(self, key, event):
        pv = self._pv[key]
        if not pv["image"]:
            return
        old = pv["scale"]
        if event.delta > 0:
            pv["scale"] = min(4.0, pv["scale"] * 1.1)
        else:
            pv["scale"] = max(0.05, pv["scale"] / 1.1)
        if old != pv["scale"]:
            self._redraw_crop(key)

    def _fill_preview_to_tree(self, key, tree):
        """将预览中选中的区域填入 Treeview 的选中行"""
        pv = self._pv[key]
        if not pv["click1"] or not pv["click2"]:
            messagebox.showinfo("提示", "请先在图片上依次点击两个点（点1 → 点2）")
            return
        sel = tree.selection()
        if not sel:
            messagebox.showinfo("提示", "请先在表格中选中一行")
            return
        x1, y1 = pv["click1"]
        x2, y2 = pv["click2"]
        tree.set(sel[0], "x1", str(x1))
        tree.set(sel[0], "y1", str(y1))
        tree.set(sel[0], "x2", str(x2))
        tree.set(sel[0], "y2", str(y2))

    def _build_ocr_tab(self, nb):
        tab = ttk.Frame(nb)
        nb.add(tab, text="数据摘取")
        frame = ttk.LabelFrame(tab, text="频点筛选")
        frame.pack(fill="x", padx=8, pady=8)
        ttk.Label(frame, text="目标频点前缀:").grid(row=0, column=0, padx=8, pady=8, sticky="e")
        ttk.Entry(frame, textvariable=self.freq_prefix, width=15).grid(row=0, column=1, padx=4, pady=8)
        ttk.Label(frame, text="前缀长度:").grid(row=1, column=0, padx=8, pady=8, sticky="e")
        ttk.Spinbox(frame, from_=1, to=20, textvariable=self.freq_length, width=8).grid(
            row=1, column=1, padx=4, pady=8, sticky="w")

        action_frame = ttk.LabelFrame(tab, text="数据提取")
        action_frame.pack(fill="x", padx=8, pady=(0, 8))
        ttk.Label(action_frame, text="从工作目录的 result.xlsx 直接提取到 finish.xlsx",
                  style="Hint.TLabel").pack(side="left", padx=8, pady=6)
        ttk.Button(action_frame, text="仅提取数据",
                   command=self._extract_only).pack(side="right", padx=8, pady=6)

        # ---- OCR 并发数（数据提取下方）----
        ocr_frame = ttk.LabelFrame(tab, text=" 并发识别")
        ocr_frame.pack(fill="x", padx=8, pady=(0, 8))
        ttk.Label(ocr_frame, text="并发引擎数:").grid(row=0, column=0, padx=8, pady=6, sticky="e")
        ttk.Spinbox(ocr_frame, from_=1, to=8, textvariable=self.ocr_concurrency,
                    width=8).grid(row=0, column=1, padx=4, pady=6, sticky="w")
        ttk.Label(ocr_frame, text="(每个引擎实例约占用 300~600MB 内存，内存小的机器建议设 1~2)",
                  style="Hint.TLabel").grid(row=0, column=2, padx=4, pady=6, sticky="w")

    def _build_image_tab(self, nb):
        tab = ttk.Frame(nb)
        nb.add(tab, text="图片插入")
        img_nb = ttk.Notebook(tab)
        img_nb.pack(fill="both", expand=True)
        self._build_img_subtab(img_nb, "主数据截图 (CQT)", self._build_cqt_main)
        self._build_img_subtab(img_nb, "Ping 截图 (CQT)", self._build_cqt_ping)
        self._build_img_subtab(img_nb, "遍历测试图 (DT)", self._build_dt)

    def _build_img_subtab(self, nb, title, builder):
        frame = ttk.Frame(nb)
        nb.add(frame, text=title)
        builder(frame)

    # ==================== 每张图位置生成 & 刷新 ====================

    def _generate_cqt_main_rows(self):
        """从每组配置 + 列配置生成每张图的行"""
        try:
            num = max(self.main_groups.get(), 0)
        except Exception:
            num = 0
        configs = self._parse_cqt_main()[:num]
        cols = self._parse_col_str(self.cqt_main_cols.get())
        rows = []
        for i, (row, hd_w, hd_h, o_w, o_h) in enumerate(configs):
            g = str(i + 1)
            for suffix, col in cols:
                w, h = (hd_w, hd_h) if suffix == 'hd' else (o_w, o_h)
                rows.append([g, f"{g}{suffix}.png", f"{col}{row}", f"{w:.4g}", f"{h:.4g}"])
        return rows

    def _generate_cqt_ping_rows(self):
        """从 ping 参数 + 列配置生成每张图的行"""
        cols = self._parse_col_str(self.ping_cols.get())
        if not cols:
            return []
        try:
            start_row = int(self.ping_start_row.get())
            row_step = int(self.ping_row_step.get())
            w = float(self.ping_w.get())
            h = float(self.ping_h.get())
            num = self.main_groups.get()
        except Exception:
            return []
        rows = []
        for i in range(1, num + 1):
            r = start_row + (i - 1) * row_step
            for suffix, col in cols:
                rows.append([str(i), f"{i}{suffix}.png", f"{col}{r}", f"{w:.4g}", f"{h:.4g}"])
        return rows

    def _generate_dt_rows(self):
        """从 DT 参数 + 列配置生成每张图的行"""
        cols = self._parse_col_str(self.dt_cols.get())
        if not cols:
            return []
        try:
            start_row = int(self.dt_start_row.get())
            row_step = int(self.dt_row_step.get())
            w = float(self.dt_w.get())
            h = float(self.dt_h.get())
            start_num = self.dt_start.get()
            end_num = self.dt_end.get()
        except Exception:
            return []
        rows = []
        for idx, folder_name in enumerate(range(start_num, end_num + 1)):
            r = start_row + idx * row_step
            for img_type, col in cols:
                rows.append([str(folder_name), f"{img_type}.png", f"{col}{r}", f"{w:.4g}", f"{h:.4g}"])
        return rows

    def _refresh_cqt_main_images(self):
        try:
            rows = self._generate_cqt_main_rows()
        except ValueError as e:
            messagebox.showwarning("提示", str(e))
            return
        self.cqt_main_img_tree.load_rows(rows)

    def _refresh_cqt_ping_images(self):
        self.cqt_ping_img_tree.load_rows(self._generate_cqt_ping_rows())

    def _refresh_dt_images(self):
        self.dt_img_tree.load_rows(self._generate_dt_rows())

    @staticmethod
    def _build_positions_from_img_tree(tree, base_dir, use_subfolder=False):
        """将每张图 treeview 行转为 {path: {position, width_cm, height_cm}}"""
        positions = {}
        for row in tree.get_all_rows():
            if len(row) < 5:
                continue
            group, img_name, cell, w_str, h_str = row
            if not img_name.strip() or not cell.strip():
                continue
            try:
                w = float(w_str)
                h = float(h_str)
            except (ValueError, TypeError):
                continue
            if use_subfolder and group.strip():
                full_path = os.path.join(base_dir, group.strip(), img_name.strip())
            else:
                full_path = os.path.join(base_dir, img_name.strip())
            positions[full_path] = {"position": cell.strip(), "width_cm": w, "height_cm": h}
        return positions

    # ==================== 图片插入子标签页 ====================

    def _build_scrollable_tree(self, parent, **kwargs):
        frame = ttk.Frame(parent)
        frame.pack(fill="both", expand=True, padx=2, pady=(2, 0))
        tree = EditableTreeview(frame, **kwargs)
        vsb = ttk.Scrollbar(frame, orient="vertical", command=tree.yview)
        hsb = ttk.Scrollbar(frame, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")
        frame.rowconfigure(0, weight=1)
        frame.columnconfigure(0, weight=1)
        return tree

    def _build_cqt_main(self, parent):
        # ---- 上半部分：字典参数 ----
        params_frame = ttk.LabelFrame(parent, text="字典参数 (每组配置)")
        params_frame.pack(fill="x", padx=4, pady=(2, 0))

        self.cqt_main_tree = EditableTreeview(
            params_frame, columns=["group","row","hd_w","hd_h","o_w","o_h"],
            col_names=["组别(自动)","起始行","hd宽cm","hd高cm","其他宽cm","其他高cm"],
            col_widths=[80,60,70,70,80,80], height=4)
        self.cqt_main_tree.pack(fill="x", padx=2, pady=(2, 0))
        for i, (row, hd_w, hd_h, o_w, o_h) in enumerate(_MAIN_SCREENSHOT_CONFIGS[:10]):
            self.cqt_main_tree.add_row([str(i+1), str(row), str(hd_w), str(hd_h), str(o_w), str(o_h)])

        bar = ttk.Frame(params_frame)
        bar.pack(fill="x", padx=4, pady=2)
        ttk.Button(bar, text="+ 添加行",
                   command=lambda: self.cqt_main_tree.add_row(
                       ["","0","10.2","12.8","8.7","12.8"])).pack(side="left")
        ttk.Button(bar, text="- 删除选中",
                   command=self.cqt_main_tree.delete_selected).pack(side="left", padx=4)

        col_frame = ttk.LabelFrame(params_frame, text="图片列配置")
        col_frame.pack(fill="x", padx=4, pady=(2, 2))
        ttk.Label(col_frame, text="格式: 后缀,列号  例: hd,AC  hu,AT  cd,BK  cu,CB").pack(side="left", padx=4)
        ttk.Entry(col_frame, textvariable=self.cqt_main_cols, width=50).pack(
            side="left", padx=4, fill="x", expand=True)

        ttk.Button(params_frame, text="刷新图片位置列表",
                   command=self._refresh_cqt_main_images).pack(pady=(0, 2))

        # ---- 下半部分：每张图位置 ----
        img_frame = ttk.LabelFrame(parent, text="图片位置详情 (每行一张图，双击编辑)")
        img_frame.pack(fill="both", expand=True, padx=4, pady=(2, 4))

        self.cqt_main_img_tree = self._build_scrollable_tree(
            img_frame, columns=["group","img_name","cell","width","height"],
            col_names=["组别","图片名","单元格","宽cm","高cm"],
            col_widths=[60,130,80,70,70])

        img_bar = ttk.Frame(img_frame)
        img_bar.pack(fill="x", padx=4, pady=2)
        ttk.Button(img_bar, text="+ 添加行",
                   command=lambda: self.cqt_main_img_tree.add_row(
                       ["","","","",""])).pack(side="left")
        ttk.Button(img_bar, text="- 删除选中",
                   command=self.cqt_main_img_tree.delete_selected).pack(side="left", padx=4)

    def _build_cqt_ping(self, parent):
        # ---- 上半部分：字典参数 ----
        params_frame = ttk.LabelFrame(parent, text="字典参数")
        params_frame.pack(fill="x", padx=4, pady=(2, 0))

        row_frame = ttk.Frame(params_frame)
        row_frame.pack(fill="x", padx=4, pady=(2, 0))
        ttk.Label(row_frame, text="起始行:").pack(side="left")
        ttk.Entry(row_frame, textvariable=self.ping_start_row, width=6).pack(side="left", padx=4)
        ttk.Label(row_frame, text="行间隔:").pack(side="left")
        ttk.Entry(row_frame, textvariable=self.ping_row_step, width=6).pack(side="left", padx=4)
        ttk.Label(row_frame, text="宽(cm):").pack(side="left")
        ttk.Entry(row_frame, textvariable=self.ping_w, width=6).pack(side="left", padx=4)
        ttk.Label(row_frame, text="高(cm):").pack(side="left")
        ttk.Entry(row_frame, textvariable=self.ping_h, width=6).pack(side="left", padx=4)

        col_frame = ttk.LabelFrame(params_frame, text="图片列配置")
        col_frame.pack(fill="x", padx=4, pady=2)
        ttk.Label(col_frame, text="格式: 后缀,列号  例: ping32,CS  ping2000,DJ").pack(side="left", padx=4)
        ttk.Entry(col_frame, textvariable=self.ping_cols, width=50).pack(
            side="left", padx=4, fill="x", expand=True)

        ttk.Button(params_frame, text="刷新图片位置列表",
                   command=self._refresh_cqt_ping_images).pack(pady=(0, 2))

        # ---- 下半部分：每张图位置 ----
        img_frame = ttk.LabelFrame(parent, text="图片位置详情 (每行一张图，双击编辑)")
        img_frame.pack(fill="both", expand=True, padx=4, pady=(2, 4))

        self.cqt_ping_img_tree = self._build_scrollable_tree(
            img_frame, columns=["group","img_name","cell","width","height"],
            col_names=["组别","图片名","单元格","宽cm","高cm"],
            col_widths=[60,130,80,70,70])

        img_bar = ttk.Frame(img_frame)
        img_bar.pack(fill="x", padx=4, pady=2)
        ttk.Button(img_bar, text="+ 添加行",
                   command=lambda: self.cqt_ping_img_tree.add_row(
                       ["","","","",""])).pack(side="left")
        ttk.Button(img_bar, text="- 删除选中",
                   command=self.cqt_ping_img_tree.delete_selected).pack(side="left", padx=4)

    def _build_dt(self, parent):
        # ---- 上半部分：字典参数 ----
        params_frame = ttk.LabelFrame(parent, text="字典参数")
        params_frame.pack(fill="x", padx=4, pady=(2, 0))

        row_frame = ttk.Frame(params_frame)
        row_frame.pack(fill="x", padx=4, pady=(2, 0))
        ttk.Label(row_frame, text="起始行:").pack(side="left")
        ttk.Entry(row_frame, textvariable=self.dt_start_row, width=6).pack(side="left", padx=4)
        ttk.Label(row_frame, text="行间隔:").pack(side="left")
        ttk.Entry(row_frame, textvariable=self.dt_row_step, width=6).pack(side="left", padx=4)
        ttk.Label(row_frame, text="宽(cm):").pack(side="left")
        ttk.Entry(row_frame, textvariable=self.dt_w, width=6).pack(side="left", padx=4)
        ttk.Label(row_frame, text="高(cm):").pack(side="left")
        ttk.Entry(row_frame, textvariable=self.dt_h, width=6).pack(side="left", padx=4)

        col_frame = ttk.LabelFrame(params_frame, text="图片列配置")
        col_frame.pack(fill="x", padx=4, pady=2)
        ttk.Label(col_frame, text="格式: 类型,列号  例: RSRP,C  SINR,I  PCI,AA  DL,O  UL,U").pack(side="left", padx=4)
        ttk.Entry(col_frame, textvariable=self.dt_cols, width=50).pack(
            side="left", padx=4, fill="x", expand=True)

        ttk.Button(params_frame, text="刷新图片位置列表",
                   command=self._refresh_dt_images).pack(pady=(0, 2))

        # ---- 下半部分：每张图位置 ----
        img_frame = ttk.LabelFrame(parent, text="图片位置详情 (每行一张图，双击编辑)")
        img_frame.pack(fill="both", expand=True, padx=4, pady=(2, 4))

        self.dt_img_tree = self._build_scrollable_tree(
            img_frame, columns=["group","img_name","cell","width","height"],
            col_names=["组别","图片名","单元格","宽cm","高cm"],
            col_widths=[60,130,80,70,70])

        img_bar = ttk.Frame(img_frame)
        img_bar.pack(fill="x", padx=4, pady=2)
        ttk.Button(img_bar, text="+ 添加行",
                   command=lambda: self.dt_img_tree.add_row(
                       ["","","","",""])).pack(side="left")
        ttk.Button(img_bar, text="- 删除选中",
                   command=self.dt_img_tree.delete_selected).pack(side="left", padx=4)

    def _build_params_tab(self, nb):
        tab = ttk.Frame(nb)
        nb.add(tab, text="运行参数")

        # ---- 模板设置 ----
        tpl_frame = ttk.LabelFrame(tab, text="Excel 模板设置")
        tpl_frame.pack(fill="x", padx=8, pady=(8, 0))

        r = 0
        ttk.Label(tpl_frame, text="模板文件:").grid(row=r, column=0, padx=8, pady=6, sticky="e")
        tpl_entry = ttk.Entry(tpl_frame, textvariable=self.template_path, width=50)
        tpl_entry.grid(row=r, column=1, padx=4, pady=6, sticky="ew")
        ttk.Button(tpl_frame, text="浏览...", command=self._browse_template).grid(
            row=r, column=2, padx=4, pady=6)
        ttk.Label(tpl_frame, text="(留空则使用工作目录下的\"室分模板.xlsx\")",
                  foreground="gray").grid(row=r, column=3, padx=4, pady=6, sticky="w")

        r += 1
        ttk.Label(tpl_frame, text="CQT 工作表名:").grid(row=r, column=0, padx=8, pady=6, sticky="e")
        ttk.Entry(tpl_frame, textvariable=self.cqt_sheet_name, width=30).grid(
            row=r, column=1, padx=4, pady=6, sticky="w")

        r += 1
        ttk.Label(tpl_frame, text="遍历测试工作表名:").grid(row=r, column=0, padx=8, pady=6, sticky="e")
        ttk.Entry(tpl_frame, textvariable=self.dt_sheet_name, width=30).grid(
            row=r, column=1, padx=4, pady=6, sticky="w")

        tpl_frame.columnconfigure(1, weight=1)

        # ---- 运行参数 ----
        frame = ttk.LabelFrame(tab, text="运行参数")
        frame.pack(fill="x", padx=8, pady=(8, 8))
        ttk.Label(frame, text="主数据/Ping 截图组数:").grid(row=0, column=0, padx=8, pady=6, sticky="e")
        ttk.Spinbox(frame, from_=1, to=50, textvariable=self.main_groups, width=8).grid(
            row=0, column=1, padx=4, pady=6, sticky="w")
        ttk.Label(frame, text="DT 起始编号:").grid(row=1, column=0, padx=8, pady=6, sticky="e")
        ttk.Spinbox(frame, from_=1, to=999, textvariable=self.dt_start, width=8).grid(
            row=1, column=1, padx=4, pady=6, sticky="w")
        ttk.Label(frame, text="DT 结束编号:").grid(row=2, column=0, padx=8, pady=6, sticky="e")
        ttk.Spinbox(frame, from_=1, to=999, textvariable=self.dt_end, width=8).grid(
            row=2, column=1, padx=4, pady=6, sticky="w")

    # ==================== 配置方案 (Profiles) ====================

    def _load_profiles(self):
        """加载所有配置方案"""
        if not os.path.exists(PROFILES_FILE):
            self._profiles = {"默认": self._collect_profile_data()}
            self._save_profiles()
        else:
            try:
                with open(PROFILES_FILE, "r", encoding="utf-8") as f:
                    self._profiles = json.load(f)
                if not isinstance(self._profiles, dict):
                    raise ValueError("配置文件格式错误")
            except (OSError, ValueError) as e:
                print(f"配置文件读取失败，使用默认配置: {e}")
                self._profiles = {"默认": self._collect_profile_data()}
                self._save_profiles()
        self._refresh_profile_list()

        # 恢复上次使用的方案
        last = self._profiles.get("__last__", "默认")
        if last in self._profiles:
            self.profile_name.set(last)
            self._apply_profile_data(self._profiles[last])

    def _save_profiles(self):
        try:
            with open(PROFILES_FILE, "w", encoding="utf-8") as f:
                json.dump(self._profiles, f, ensure_ascii=False, indent=2)
        except OSError as e:
            print(f"配置文件保存失败: {e}")

    def _refresh_profile_list(self):
        names = [k for k in self._profiles if not k.startswith("__")]
        self.profile_combo["values"] = names

    def _collect_profile_data(self):
        """收集当前所有界面参数为一个字典"""
        return {
            "base_dir": self.base_dir.get(),
            "template_path": self.template_path.get(),
            "cqt_sheet_name": self.cqt_sheet_name.get(),
            "dt_sheet_name": self.dt_sheet_name.get(),
            "freq_prefix": self.freq_prefix.get(),
            "freq_length": self.freq_length.get(),
            "ocr_concurrency": self.ocr_concurrency.get(),
            "main_groups": self.main_groups.get(),
            "dt_start": self.dt_start.get(),
            "dt_end": self.dt_end.get(),
            "ping_regions": self._ping_tree.get_all_rows(),
            "main_regions": self._main_tree.get_all_rows(),
            "cqt_main_configs": self.cqt_main_tree.get_all_rows(),
            "cqt_main_cols": self.cqt_main_cols.get(),
            "ping_start_row": self.ping_start_row.get(),
            "ping_row_step": self.ping_row_step.get(),
            "ping_w": self.ping_w.get(),
            "ping_h": self.ping_h.get(),
            "ping_cols": self.ping_cols.get(),
            "dt_start_row": self.dt_start_row.get(),
            "dt_row_step": self.dt_row_step.get(),
            "dt_w": self.dt_w.get(),
            "dt_h": self.dt_h.get(),
            "dt_cols": self.dt_cols.get(),
            "cqt_main_images": self.cqt_main_img_tree.get_all_rows(),
            "cqt_ping_images": self.cqt_ping_img_tree.get_all_rows(),
            "dt_images": self.dt_img_tree.get_all_rows(),
        }

    def _apply_profile_data(self, d):
        """将字典数据恢复到界面"""
        self.base_dir.set(d.get("base_dir", self.base_dir.get()))
        self.template_path.set(d.get("template_path", ""))
        self.cqt_sheet_name.set(d.get("cqt_sheet_name", "SA-网络性能验收-CQT"))
        self.dt_sheet_name.set(d.get("dt_sheet_name", "SA-网络性能验收-遍历测试图"))
        self.freq_prefix.set(d.get("freq_prefix", "525150"))
        self.freq_length.set(d.get("freq_length", 6))
        self.ocr_concurrency.set(d.get("ocr_concurrency", _OCR_CONCURRENCY))
        self.main_groups.set(d.get("main_groups", 4))
        self.dt_start.set(d.get("dt_start", 1))
        self.dt_end.set(d.get("dt_end", 6))
        self._ping_tree.load_rows(d.get("ping_regions", []))
        self._main_tree.load_rows(d.get("main_regions", []))
        self.cqt_main_tree.load_rows(d.get("cqt_main_configs", []))
        self.cqt_main_cols.set(d.get("cqt_main_cols", "hd,AC hu,AT cd,BK cu,CB"))
        self.ping_start_row.set(d.get("ping_start_row", "6"))
        self.ping_row_step.set(d.get("ping_row_step", "22"))
        self.ping_w.set(d.get("ping_w", "8.6"))
        self.ping_h.set(d.get("ping_h", "8.5"))
        self.ping_cols.set(d.get("ping_cols", "ping32,CS ping2000,DJ"))
        self.dt_start_row.set(d.get("dt_start_row", "7"))
        self.dt_row_step.set(d.get("dt_row_step", "18"))
        self.dt_w.set(d.get("dt_w", "10.2"))
        self.dt_h.set(d.get("dt_h", "7.79"))
        self.dt_cols.set(d.get("dt_cols", "RSRP,C SINR,I PCI,AA DL,O UL,U"))
        self.cqt_main_img_tree.load_rows(d.get("cqt_main_images", []))
        self.cqt_ping_img_tree.load_rows(d.get("cqt_ping_images", []))
        self.dt_img_tree.load_rows(d.get("dt_images", []))

    def _on_profile_select(self, event=None):
        name = self.profile_name.get()
        if name in self._profiles:
            self._apply_profile_data(self._profiles[name])

    def _save_profile(self):
        name = self.profile_name.get()
        self._profiles[name] = self._collect_profile_data()
        self._profiles["__last__"] = name
        self._save_profiles()
        self._refresh_profile_list()
        self.status_var.set(f"方案「{name}」已保存")

    def _save_as_profile(self):
        name = simpledialog.askstring("另存为", "输入配置方案名称:", parent=self)
        if name:
            self.profile_name.set(name)
            self._save_profile()

    def _delete_profile(self):
        name = self.profile_name.get()
        if name == "默认":
            messagebox.showwarning("提示", "不能删除「默认」方案")
            return
        if messagebox.askyesno("确认", f"删除配置方案「{name}」?"):
            del self._profiles[name]
            if self._profiles.get("__last__") == name:
                self._profiles["__last__"] = "默认"
            self._save_profiles()
            self.profile_name.set("默认")
            self._refresh_profile_list()
            self._apply_profile_data(self._profiles["默认"])
            self.status_var.set(f"方案「{name}」已删除")

    # ==================== 关闭 ====================

    def _on_close(self):
        self._profiles["__last__"] = self.profile_name.get()
        self._save_profiles()
        self.destroy()

    # ==================== 运行 ====================

    def _browse_dir(self):
        path = filedialog.askdirectory(title="选择工作目录")
        if path:
            self.base_dir.set(path)

    def _browse_template(self):
        path = filedialog.askopenfilename(
            title="选择 Excel 模板文件",
            filetypes=[("Excel 文件", "*.xlsx"), ("所有文件", "*.*")])
        if path:
            self.template_path.set(path)

    def _start_run(self):
        if self._running:
            return
        self._cancel.clear()
        self._run_error = False
        self._error_shown = False
        self._running = True
        self.start_btn.config(state="disabled")
        self.stop_btn.config(state="normal")
        self.progress["value"] = 0
        self.status_var.set("0%")
        self._set_status_dot(self._colors["running"])
        threading.Thread(target=self._run, daemon=True).start()

    def _stop_run(self):
        if self._running and not self._cancel.is_set():
            self._cancel.set()
            self.status_var.set("正在停止（当前步骤完成后生效）...")

    def _extract_only(self):
        if self._running:
            return
        self._cancel.clear()
        self._run_error = False
        self._error_shown = False
        self._running = True
        self.start_btn.config(state="disabled")
        self.stop_btn.config(state="normal")
        self.progress["value"] = 0
        self.status_var.set("0%")
        self._set_status_dot(self._colors["running"])
        threading.Thread(target=self._extract_only_worker, daemon=True).start()

    def _extract_only_worker(self):
        try:
            base = self.base_dir.get().strip()
            if not base:
                raise ValueError("请先设置工作目录")
            result_excel = os.path.join(base, "result.xlsx")
            finish_excel = os.path.join(base, "finish.xlsx")
            try:
                freq_prefix = int(self.freq_prefix.get())
            except ValueError:
                raise ValueError("目标频点前缀必须是数字")
            freq_len = self.freq_length.get()
            if freq_len < 1 or freq_len > 20:
                raise ValueError("前缀长度必须在 1 到 20 之间")
            self._post_progress(30, "正在读取 result.xlsx...")
            extract_result_to_finish(result_excel, finish_excel,
                                     target_freq_prefix=freq_prefix,
                                     prefix_length=freq_len)
            self._post_progress(100, "提取完成")
        except Exception as e:
            import traceback
            traceback.print_exc()
            self._post_progress(0, f"提取失败: {e}")
        finally:
            self._post_done()

    def _post_progress(self, pct, step_text=""):
        self._ui_queue.put((pct, step_text))

    def _post_done(self):
        self._ui_queue.put("done")

    def _poll_ui_queue(self):
        try:
            while True:
                item = self._ui_queue.get_nowait()
                if item == "done":
                    self._running = False
                    self.start_btn.config(state="normal")
                    self.stop_btn.config(state="disabled")
                    if not self._run_error:
                        self._set_status_dot(self._colors["success"])
                else:
                    pct, text = item
                    self._update_progress(pct, text)
        except queue.Empty:
            pass
        self.after(100, self._poll_ui_queue)

    def _update_progress(self, pct, step_text=""):
        self.progress["value"] = pct
        is_error = bool(step_text) and ("错误" in step_text or "失败" in step_text)
        if is_error:
            self.status_var.set(f"{pct}%  {step_text}")
        else:
            self.status_var.set(f"{pct}%")
        if is_error:
            self._run_error = True
            self._set_status_dot(self._colors["error"])
            # 失败弹窗（只弹一次，避免多个错误重复打扰）
            if not self._error_shown:
                self._error_shown = True
                self.after(50, lambda: messagebox.showerror("运行失败", step_text))
        elif pct >= 100:
            self._run_error = False
            self._set_status_dot(self._colors["success"])
        else:
            self._set_status_dot(self._colors["running"])

    def _run(self):
        try:
            base = self.base_dir.get()
            # 模板路径：优先使用用户选择的路径，否则使用工作目录下的默认文件名
            tpl = self.template_path.get().strip()
            if tpl:
                template_excel = tpl
            else:
                template_excel = os.path.join(base, "室分模板.xlsx")

            # 预检
            if not os.path.exists(template_excel):
                self._post_progress(0, "错误: 模板文件不存在")
                return
            try:
                wb = load_workbook(template_excel)
                wb.close()
            except Exception as e:
                self._post_progress(0, "错误: 模板文件无效 (不是 .xlsx 格式)")
                return

            # 解析参数
            ping_regions = self._parse_regions(self._ping_tree)
            main_regions = self._parse_regions(self._main_tree)
            cqt_main_configs = self._parse_cqt_main()
            cqt_main_cols = self._parse_col_str(self.cqt_main_cols.get())
            ping_cols = self._parse_col_str(self.ping_cols.get())
            dt_cols = self._parse_col_str(self.dt_cols.get())
            dt_types = [c[0] for c in dt_cols]
            dt_columns = [c[1] for c in dt_cols]
            num_groups = self.main_groups.get()
            if num_groups < 1:
                raise ValueError("主数据/Ping 截图组数必须 ≥ 1")
            try:
                freq_prefix = int(self.freq_prefix.get())
            except ValueError:
                raise ValueError("目标频点前缀必须是数字")
            freq_len = self.freq_length.get()
            if freq_len < 1 or freq_len > 20:
                raise ValueError("前缀长度必须在 1 到 20 之间")
            concurrency = self.ocr_concurrency.get()
            if concurrency < 1 or concurrency > 8:
                raise ValueError("OCR 并发引擎数必须在 1 到 8 之间")
            if self.dt_start.get() > self.dt_end.get():
                raise ValueError("DT 起始编号不能大于结束编号")
            cqt_sheet = self.cqt_sheet_name.get().strip() or 'SA-网络性能验收-CQT'
            dt_sheet = self.dt_sheet_name.get().strip() or 'SA-网络性能验收-遍历测试图'

            # 从每张图 treeview 构建 positions（有数据则覆盖参数计算）
            pictuer_dir = os.path.join(base, "pictuer")
            ping_dir = os.path.join(base, "ping")
            dt_dir = os.path.join(base, "dt")

            cqt_main_override = None
            if hasattr(self, 'cqt_main_img_tree') and self.cqt_main_img_tree.get_children():
                positions = self._build_positions_from_img_tree(
                    self.cqt_main_img_tree, pictuer_dir, use_subfolder=False)
                if positions:
                    cqt_main_override = positions

            cqt_ping_override = None
            if hasattr(self, 'cqt_ping_img_tree') and self.cqt_ping_img_tree.get_children():
                positions = self._build_positions_from_img_tree(
                    self.cqt_ping_img_tree, ping_dir, use_subfolder=False)
                if positions:
                    cqt_ping_override = positions

            dt_override = None
            if hasattr(self, 'dt_img_tree') and self.dt_img_tree.get_children():
                positions = self._build_positions_from_img_tree(
                    self.dt_img_tree, dt_dir, use_subfolder=True)
                if positions:
                    dt_override = positions

            steps = [
                ("第一步: 裁剪 ping 截图",
                 lambda: step1_crop_ping(
                     os.path.join(base, "ping"), os.path.join(base, "jietu"),
                     regions=ping_regions)),
                ("第二步: 裁剪主数据截图",
                 lambda: step2_crop_main(
                     os.path.join(base, "pictuer"), os.path.join(base, "jietu"),
                     regions=main_regions)),
                ("第三步: OCR 识别 + 数据提取",
                 lambda: step3_ocr_and_extract(
                     os.path.join(base, "jietu"), os.path.join(base, "result.xlsx"),
                     os.path.join(base, "finish.xlsx"),
                     target_freq_prefix=freq_prefix, prefix_length=freq_len,
                     concurrency=concurrency)),
                ("第四步: 插入主数据截图",
                 lambda: step4_insert_main_screenshots(
                     template_excel, num_groups=num_groups,
                     configs=cqt_main_configs, img_columns=cqt_main_cols,
                     pictuer_dir=pictuer_dir,
                     sheet_name=cqt_sheet,
                     positions_override=cqt_main_override)),
                ("第五步: 插入 ping 截图",
                 lambda: step5_insert_ping_screenshots(
                     template_excel, num_groups=num_groups,
                     start_row=int(self.ping_start_row.get()),
                     row_step=int(self.ping_row_step.get()),
                     img_columns=ping_cols,
                     width_cm=float(self.ping_w.get()),
                     height_cm=float(self.ping_h.get()),
                     ping_dir=ping_dir,
                     sheet_name=cqt_sheet,
                     positions_override=cqt_ping_override)),
                ("第六步: 插入遍历测试图",
                 lambda: step6_insert_dt_screenshots(
                     template_excel,
                     start_num=self.dt_start.get(), end_num=self.dt_end.get(),
                     img_types=dt_types, img_cols=dt_columns,
                     start_row=int(self.dt_start_row.get()),
                     row_step=int(self.dt_row_step.get()),
                     width_cm=float(self.dt_w.get()),
                     height_cm=float(self.dt_h.get()),
                     dt_dir=dt_dir,
                     sheet_name=dt_sheet,
                     positions_override=dt_override)),
            ]

            for i, (desc, fn) in enumerate(steps):
                if self._cancel.is_set():
                    self._post_progress(0, "已停止")
                    break
                self._post_progress((i + 1) * 100 // len(steps), desc)
                fn()

            self._post_progress(100, "全部完成!")
            print("===== 全部完成! =====")  # 仅输出到终端

        except Exception as e:
            import traceback
            traceback.print_exc()
            self._post_progress(0, f"运行失败: {e}")
        finally:
            self._post_done()

    # ==================== 解析辅助 ====================

    def _parse_regions(self, tree):
        regions = []
        for idx, row in enumerate(tree.get_all_rows(), 1):
            if len(row) < 5:
                continue
            name = (row[0] or "").strip()
            if not name:
                continue
            try:
                x1, y1, x2, y2 = (int(row[i]) for i in range(1, 5))
            except (ValueError, TypeError):
                raise ValueError(f"裁剪区域第 {idx} 行（{name}）的坐标必须是整数")
            regions.append((name, x1, y1, x2, y2))
        return regions

    def _parse_cqt_main(self):
        configs = []
        for idx, row in enumerate(self.cqt_main_tree.get_all_rows(), 1):
            if len(row) < 6:
                continue
            try:
                configs.append((int(row[1]), float(row[2]), float(row[3]),
                                float(row[4]), float(row[5])))
            except (ValueError, TypeError):
                raise ValueError(f"CQT 配置第 {idx} 行的起始行/宽/高必须是数字")
        return configs

    @staticmethod
    def _parse_col_str(s):
        result = []
        for part in s.strip().split():
            if ',' in part:
                a, b = part.split(',', 1)
                result.append((a.strip(), b.strip()))
        return result


# ======================== 主入口 ========================

def _enable_high_dpi():
    """启用高 DPI 清晰显示（Windows 高分屏 125%/150% 缩放不模糊）。"""
    if sys.platform == "win32":
        try:
            import ctypes
            try:
                # Windows 10 1703+：按监视器 DPI 感知
                ctypes.windll.shcore.SetProcessDpiAwareness(2)
            except Exception:
                ctypes.windll.user32.SetProcessDPIAware()
        except Exception:
            pass


if __name__ == "__main__":
    _enable_high_dpi()
    check_license()
    app = App()
    app.mainloop()
